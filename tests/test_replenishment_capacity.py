from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import Mock, MagicMock, patch

import pymysql
import pytest
from openpyxl import load_workbook

from app import main
from app.services.replenishment_capacity import (
    CAPACITY_COLUMNS, FT3_TO_M3, ReplenishmentCapacityService, dimensions_m3, evaluate_capacity, volume_m3,
)

DAY = date(2026, 10, 14)


def test_capacity_sources_are_scanned_once_for_many_distinct_products():
    class CountedRows(list):
        scans = 0
        def __iter__(self):
            self.scans += 1
            return super().__iter__()

    evidence = sources()
    source_keys = ("fees", "warehouse_types", "listing_dimensions", "packages", "limits", "sellers")
    for key in source_keys:
        evidence[key] = CountedRows(evidence.get(key, []))
    rows = [dict(country="欧洲站", store="shop", msku=f"other-{i}", replenish_qty=0) for i in range(40)]
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(return_value=evidence)
    service.enrich(rows, DAY, demand_rows=[])
    assert all(evidence[key].scans <= 1 for key in source_keys)
    assert all(row["capacity_status"] == "no_replenishment" for row in rows)


def test_capacity_lookup_does_not_rescan_large_tables_per_200_products():
    conn, cursor = MagicMock(), MagicMock()
    conn.__enter__.return_value = conn
    conn.cursor.return_value.__enter__.return_value = cursor
    active_sql = [""]
    cursor.execute.side_effect = lambda sql, args=(): active_sql.__setitem__(0, sql)
    cursor.fetchall.side_effect = lambda: [{"report_month":"2026-09"}] if "max(month_of_charge)" in active_sql[0] else []
    rows = [dict(store="shop", msku=f"msku-{i}", sku=f"sku-{i}") for i in range(401)]
    ReplenishmentCapacityService(lambda: conn)._load(rows, DAY)
    for table in ("lx_statistics_storage_fee_month", "lx_sales_mws_listing", "FBA库存明细"):
        queries = [call for call in cursor.execute.call_args_list if table in call.args[0] and " in (" in call.args[0]]
        assert len(queries) == 1


def test_excel_export_reuses_alignment_styles_instead_of_allocating_per_cell():
    payload = {"columns":[{"name":"seller_sku_adj", "label":"MSKU"},
                           {"name":"capacity_note", "label":"库容说明"}],
               "rows":[{"seller_sku_adj":f"m-{i}", "capacity_note":"本次补货未超库容"} for i in range(100)]}
    with patch.object(main, "Alignment", wraps=main.Alignment) as factory:
        data = main.build_replenishment_xlsx(payload)
    assert factory.call_count < 200  # The fixed legend has its own styles; data rows add none.
    book = load_workbook(BytesIO(data))
    assert book.active["A2"].alignment.vertical == "center"
    assert book.active["B2"].alignment.vertical == "top" and book.active["B2"].alignment.wrap_text
    assert book.active["C2"].alignment.vertical == "center" and book.active["C2"].alignment.wrap_text


def test_independent_capacity_queries_overlap_but_keep_original_result_order():
    from threading import Barrier, Lock
    import time
    service = ReplenishmentCapacityService(Mock())
    barrier, lock = Barrier(3), Lock()
    state = {"active":0, "peak":0}
    def read(job):
        with lock:
            state["active"] += 1
            state["peak"] = max(state["peak"], state["active"])
        barrier.wait(timeout=3)
        time.sleep((3-job[2])*0.01)
        with lock:
            state["active"] -= 1
        return job[0], [{"order":job[2]}]
    service._query_source = read
    source = {"fees":[]}
    service._load_source_queries(source, [("fees","sql",i) for i in range(3)])
    assert state["peak"] == 3
    assert source["fees"] == [{"order":0},{"order":1},{"order":2}]


def context(**updates):
    return dict(group="shop / EU", group_confirmed=True, type_label="标准尺寸", type_confirmed=True,
                month="2026-10-01", source_time=datetime(2026,10,14,8), maximum="10", remaining="2", unit=1,
                allocation_type=2, unit_volume=Decimal(".05"), measurement_source="模拟的明确仓储类型和尺寸",
                demand_month="2026-10", demand_complete=True, occupancy_reconciled=True,
                same_month_new_volume=Decimal("1"), **updates)


@pytest.mark.parametrize("remaining,demand,status", [
    ("-1","1","over_capacity"), ("0","1","no_capacity"),
    ("2","3","insufficient"), ("2","2","within_capacity"), ("2","1","within_capacity"),
])
def test_confirmed_capacity_boundaries(remaining, demand, status):
    source = context()
    source.update(remaining=remaining, same_month_new_volume=demand)
    result = evaluate_capacity(DAY,20,source)
    assert result["capacity_status"] == status
    assert result["capacity_demand_volume_m3"] == 1
    if status in {"within_capacity","insufficient"}:
        assert result["capacity_shortfall_m3"] == max(0,float(demand)-float(remaining))


@pytest.mark.parametrize("changes,reason,status", [
    ({"maximum":"-","remaining":"-1"}, "额度或单位", "pending"),
    ({"unit":99}, "额度或单位", "pending"), ({"unit_volume":None}, "单件体积", "source_available"),
    ({"type_confirmed":False}, "仓储类型需确认", "source_available"), ({"group_confirmed":False}, "共享库容组", "source_available"),
    ({"source_time":"2026-10-12 08:00:00"}, "来源未覆盖", "pending"),
    ({"source_time":"2026-10-15 08:00:00"}, "来源晚于", "pending"),
    ({"month":"2026-11-01"}, "额度月份", "pending"), ({"allocation_type":1}, "预估", "pending"),
    ({"demand_month":"2026-11"}, "发货月份", "source_available"), ({"demand_complete":False}, "未完整", "source_available"),
    ({"occupancy_reconciled":False}, "货件占用尚未核对", "source_available"),
])
def test_uncertain_evidence_never_claims_replenishment_fits_but_preserves_source_quota(changes,reason,status):
    source = context()
    source.update(changes)
    result = evaluate_capacity(DAY,20,source)
    assert result["capacity_status"] == status
    assert reason in result["capacity_note"]
    assert result["capacity_group_demand_m3"] is None
    assert result["capacity_shortfall_m3"] is None
    if changes.get("maximum") == "-":
        assert result["capacity_remaining_m3"] is None


def test_old_negative_capacity_remains_an_observation_not_current_proof():
    source = context()
    source.update(remaining="-.001", source_time="2026-10-12 08:00:00", demand_month=None)
    result = evaluate_capacity(DAY,20,source)
    assert result["capacity_status"] == "pending"
    assert result["capacity_status_label"] == "库容待核（来源超占用）"
    assert result["capacity_remaining_m3"] == -.001
    assert "2026-10-12" in result["capacity_detail"]
    assert "默认发货月份：2026-10" in result["capacity_note"]


@pytest.mark.parametrize("source_day,expected", [("2026-10-14", "within_capacity"),
                                                ("2026-10-13", "within_capacity"),
                                                ("2026-10-12", "pending"),
                                                ("2026-10-15", "pending")])
def test_daily_capacity_allows_previous_day_snapshot(source_day, expected):
    evidence = context()
    evidence.update(source_time=source_day, demand_month=None)
    result = evaluate_capacity(DAY, 20, evidence)
    assert result["capacity_status"] == expected
    assert "发货月份未确定" not in result["capacity_note"]
    assert "默认发货月份：2026-10" in result["capacity_note"]


def test_snapshot_freshness_uses_collection_date_and_preserves_platform_update_time():
    evidence = context()
    evidence.update(source_time="2026-10-09 08:00:00", snapshot_time="2026-10-13")
    result = evaluate_capacity(DAY, 20, evidence)
    assert result["capacity_status"] == "within_capacity"
    assert result["capacity_source_time"] == "2026-10-09 08:00:00"
    assert "库容快照日期：2026-10-13" in result["capacity_note"]


def test_month_boundary_defaults_shipping_to_generation_month():
    evidence = context()
    evidence.update(month="2026-11-01", source_time="2026-10-31", demand_month=None)
    result = evaluate_capacity(date(2026, 11, 1), 20, evidence)
    assert result["capacity_status"] == "within_capacity"
    assert "默认发货月份：2026-11" in result["capacity_note"]


def test_confirming_shipping_month_does_not_confirm_product_type_or_double_count_cargo():
    evidence = context()
    evidence.update(type_confirmed=False, demand_month=None, source_time="2026-10-13")
    result = evaluate_capacity(DAY, 20, evidence)
    assert result["capacity_status"] == "source_available"
    assert "仓储类型需确认" in result["capacity_note"]
    assert "发货月份未确定" not in result["capacity_note"]
    evidence.update(type_confirmed=True, occupancy_reconciled=False)
    result = evaluate_capacity(DAY, 20, evidence)
    assert result["capacity_status"] == "source_available"
    assert "货件占用尚未核对" in result["capacity_note"]


def test_units_and_amazon_dimension_precision():
    assert volume_m3("100",2) == Decimal(100)*FT3_TO_M3
    assert volume_m3("-1",2) < 0
    assert volume_m3("-",1) is None
    source = dict(longest_side="28",median_side="19",shortest_side="7",item_volume="0.00")
    assert dimensions_m3(source,("longest_side","median_side","shortest_side"),"centimeters") == Decimal(".003724")
    assert dimensions_m3(source,("longest_side","median_side","shortest_side"),"inches") == Decimal(28)*19*7*Decimal(".0254")**3
    source["shortest_side"] = 0
    assert dimensions_m3(source,("longest_side","median_side","shortest_side"),"centimeters") is None


def sources():
    return dict(sellers=[dict(sid=1,name="shop-DE",account_name="shop",seller_account_id=7,country="德国"),
                         dict(sid=2,name="shop-UK",account_name="shop",seller_account_id=7,country="英国")],
                groups=[dict(seller_account_id=7,account_name="shop",capacity_group="EU",representative_sid=1,
                             capacity_source_update_time=datetime(2026,10,14,8),remaining_volume_m3=Decimal("-1")),
                        dict(seller_account_id=7,account_name="shop",capacity_group="UK",representative_sid=2,
                             capacity_source_update_time=datetime(2026,10,14,8),remaining_volume_m3=Decimal("999"))],
                limits=[dict(sid="1",month="2026-10-01",sub_item_storage_type="Standard",sub_item_vol_unit_type=1,
                             vol_stock_max="10",vol_stock_remain="-1",source_update_time=datetime(2026,10,14,8),fba_allocation_type=2),
                        dict(sid="1",month="2026-10-01",sub_item_storage_type="Oversize",sub_item_vol_unit_type=2,
                             vol_stock_max="200",vol_stock_remain="100",source_update_time=datetime(2026,10,14,8),fba_allocation_type=2)],
                fees=[dict(seller_name="shop-DE",country="德国",seller_sku="msku",product_size_tier="Standard-Size",
                           longest_side=10,median_side=10,shortest_side=10,measurement_units="centimeters",month_of_charge="2026-09"),
                      dict(seller_name="shop-UK",country="英国",seller_sku="msku",product_size_tier="Oversize",
                           longest_side=100,median_side=100,shortest_side=100,measurement_units="centimeters",month_of_charge="2026-09")],packages=[])


def test_country_account_type_isolation_and_no_standard_group_reuse_for_oversize():
    service = ReplenishmentCapacityService(Mock())
    source = sources()
    row = dict(country_category="欧洲站",seller_name_new="shop",seller_sku_adj="msku",replenish_qty=20)
    result = service._context(row,source)
    assert result["remaining"] == "-1"
    assert result["unit_volume"] == Decimal(".001")
    assert result["type_confirmed"] is False
    source["fees"][0]["product_size_tier"] = "Oversize"
    result = service._context(row,source)
    assert result["remaining"] == "100" and result["unit"] == 2
    assert volume_m3(result["remaining"],result["unit"]) == Decimal(100)*FT3_TO_M3
    row["seller_name_new"] = "other-account"
    result = service._context(row,source)
    assert not result["group_confirmed"]
    assert result.get("unit_volume") is None


def test_production_context_always_uses_store_dwd_collection_time_and_platform_time():
    service = ReplenishmentCapacityService(Mock())
    evidence = sources()
    evidence["snapshot_time"] = date(2026, 10, 13)
    evidence["groups"][0]["capacity_source_update_time"] = datetime(2026, 10, 9, 8)
    evidence["limits"][0]["create_time"] = datetime(2026, 10, 13, 9)
    evidence["limits"][0]["source_update_time"] = datetime(2026, 10, 10, 8)
    evidence["limits"][1]["create_time"] = datetime(2026, 10, 13, 9)
    row = dict(country_category="欧洲站", seller_name_new="shop", seller_sku_adj="msku")
    standard = service._context(row, evidence)
    assert standard["snapshot_time"] == datetime(2026, 10, 13, 9)
    assert standard["source_time"] == datetime(2026, 10, 10, 8)
    evidence["fees"][0]["product_size_tier"] = "Oversize"
    oversize = service._context(row, evidence)
    assert oversize["snapshot_time"] == datetime(2026, 10, 13, 9)
    assert oversize["unit"] == 2  # Retains the type's own quota and unit.


def test_page_and_export_shapes_have_identical_annotations_and_never_modify_original_fields():
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(return_value=sources())
    raw = dict(country_category="欧洲站",seller_name_new="shop",seller_sku_adj="msku",max_sku="local",replenish_qty=20,
               replenish_cost=300,replenish_box_qty=2,abcd_category="明星产品")
    page = dict(country="欧洲站",store="shop",msku="msku",sku="local",replenish_qty=20,cost=300,box_qty=2,category="明星产品")
    saved = deepcopy(raw)
    service.enrich([raw],DAY)
    service.enrich([page],DAY)
    assert {key:raw[key] for key in saved} == saved
    for column in CAPACITY_COLUMNS:
        assert page[column["name"]] == raw[column["name"]]
    service.enrich([raw,deepcopy(raw)],DAY)
    assert raw["capacity_remaining_m3"] == -1  # Not summed or consumed per product.
    assert raw["capacity_group_demand_m3"] is None


def test_source_failure_is_visible_and_does_not_break_replenishment():
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(side_effect=pymysql.OperationalError(2003,"unavailable"))
    row = dict(store="shop",msku="msku",replenish_qty=20,cost=300)
    service.enrich([row],DAY)
    assert row["capacity_status"] == "pending" and "暂不可读取" in row["capacity_note"]
    assert row["replenish_qty"] == 20 and row["cost"] == 300


def test_export_capacity_columns_and_legend_preserve_warning_fill_and_follow_frame():
    row = dict(seller_sku_adj="msku",replenish_qty=20,support_replenish_level_sort=2,
               final_sales_3d=20,final_sales_7d=20,fllow_flag=0)
    row.update(evaluate_capacity(DAY,20,context()))
    columns = [{"name":"seller_sku_adj","label":"MSKU"},{"name":"replenish_qty","label":"补货数量"},*CAPACITY_COLUMNS]
    workbook = load_workbook(BytesIO(main.build_replenishment_xlsx({"columns":columns,"rows":[row]})))
    sheet = workbook.active
    assert sheet.cell(2,2).value == 20 and sheet.max_row == 2
    assert all(cell.fill.fgColor.rgb.endswith("FFF2CC") for cell in sheet[2])
    assert sheet.cell(2,sheet.max_column).border.right.style == "thick"
    assert "同期需求未超额" in sheet.cell(2,sheet.max_column).value
    legend = "\n".join(str(cell.value or "") for cells in workbook.worksheets[-1] for cell in cells)
    for label in ["库容已超占用/无剩余", "预计库容不足", "同期需求未超额", "库容待核", "四种销售角色"]:
        assert label in legend
    assert "前一天" in legend and "当月" in legend
    assert "发货月份未知" not in legend


def test_store_quota_is_loaded_without_shared_group():
    service = ReplenishmentCapacityService(Mock())
    evidence = sources()
    evidence["groups"] = []
    row = dict(country_category="欧洲站", seller_name_new="shop", seller_sku_adj="msku", replenish_qty=20)
    source = service._context(row, evidence)
    assert source["remaining"] == "-1"
    result = evaluate_capacity(DAY, 20, source)
    assert result["capacity_status"] == "over_capacity"
    assert result["capacity_remaining_m3"] == -1
    assert "sid=1" in result["capacity_match_basis"]
    assert "商品仓储类型" in result["capacity_product_note"]


def test_no_batch_does_not_prevent_loading_store_quota():
    conn, cursor = MagicMock(), MagicMock()
    conn.__enter__.return_value = conn
    conn.cursor.return_value.__enter__.return_value = cursor
    fixtures = sources()
    active_sql = [""]
    cursor.execute.side_effect = lambda sql, args=(): active_sql.__setitem__(0, sql)
    def records():
        sql = active_sql[0]
        if "lx_basic_seller" in sql:
            return fixtures["sellers"]
        if "lx_replenishment_restriction" in sql:
            return fixtures["limits"]
        if "max(month_of_charge)" in sql:
            return [{"report_month": "2026-09"}]
        if "lx_statistics_storage_fee_month" in sql:
            return fixtures["fees"]
        return []
    cursor.fetchall.side_effect = records
    result = ReplenishmentCapacityService(lambda: conn)._load([dict(country="欧洲站", store="shop", msku="msku")], DAY)
    assert result.get("limits") == fixtures["limits"]
    assert result.get("groups") == []
    assert not result.get("error")
    listing_sql = next(call.args[0] for call in cursor.execute.call_args_list if "dimension_info" in call.args[0])
    assert "partition by sid, binary seller_sku" in listing_sql
    assert "create_time=(select max" not in listing_sql  # Batch timestamps differ by product/site.
    # Avoid materializing the large JSON payload in the window-function sort.
    ranked_sql = listing_sql.split("from (", 1)[1].split(") latest_dimensions", 1)[0]
    assert "dimension_info" not in ranked_sql
    assert "listing.id=latest_dimensions.id" in listing_sql
    listing_args = next(call.args[1] for call in cursor.execute.call_args_list if "dimension_info" in call.args[0])
    assert "and binary seller_sku in" in listing_sql
    assert "1" in listing_args and "2" not in listing_args  # Only this object's European site ids.


def test_existing_quota_is_visible_when_product_size_type_is_unknown():
    service = ReplenishmentCapacityService(Mock())
    evidence = sources()
    evidence.update(groups=[], fees=[])
    result = evaluate_capacity(DAY, 20, service._context(dict(country="欧洲站", store="shop", msku="missing"), evidence))
    assert result["capacity_status"] == "type_unmapped"
    assert "标准尺寸" in result["capacity_quota_summary"] and "大件" in result["capacity_quota_summary"]
    assert result["capacity_remaining_m3"] is None  # Different types are never summed.
    assert result["capacity_month"] == "2026-10-01"


def test_store_station_disagreement_is_listed_without_adding_or_borrowing_capacity():
    service = ReplenishmentCapacityService(Mock())
    evidence = sources()
    evidence["groups"] = []
    evidence["sellers"].append(dict(sid=3, name="shop-FR", account_name="shop", seller_account_id=7, country="法国"))
    other = deepcopy(evidence["limits"][0])
    other.update(sid="3", vol_stock_remain="2")
    evidence["limits"].append(other)
    result = evaluate_capacity(DAY, 20, service._context(dict(country="欧洲站", store="shop", msku="msku"), evidence))
    assert result["capacity_status"] == "station_difference"
    assert result["capacity_remaining_m3"] is None
    assert "德国" in result["capacity_quota_summary"] and "法国" in result["capacity_quota_summary"]


def test_confirmed_group_does_not_supply_another_stores_representative_sid():
    service = ReplenishmentCapacityService(Mock())
    evidence = sources()
    evidence["groups"][0].update(representative_sid=99, remaining_volume_m3=Decimal("9999"))
    result = service._context(dict(country="欧洲站", store="shop", msku="msku"), evidence)
    assert result["remaining"] == "-1"
    assert "sid=1" in result["match_basis"]


def test_historical_month_uses_month_end_archive_with_explicit_label():
    evidence = context()
    evidence.update(month="2026-09-01", source_time="2026-09-30", snapshot_time="2026-09-30",
                    reference_date=date(2026, 9, 30), archive_mode=True, demand_month="2026-09")
    result = evaluate_capacity(date(2026, 9, 10), 20, evidence)
    assert result["capacity_status"] == "within_capacity"
    assert "月末归档" in result["capacity_note"]
    assert "不代表所选历史日" in result["capacity_note"]


def test_missing_measurements_do_not_hide_valid_positive_store_quota():
    evidence = context()
    evidence.update(type_confirmed=False, group_confirmed=False, unit_volume=None)
    result = evaluate_capacity(DAY, 20, evidence)
    assert result["capacity_status"] == "source_available"
    assert result["capacity_remaining_m3"] == 2
    assert result["capacity_shortfall_m3"] is None
    assert "单件体积" in result["capacity_product_note"]


def test_equal_station_quotas_are_normalized_but_not_summed_or_confirmed_shared():
    service = ReplenishmentCapacityService(Mock())
    evidence = sources()
    evidence["groups"] = []
    evidence["sellers"].append(dict(sid=3, name="shop-FR", account_name="shop", seller_account_id=7, country="法国"))
    evidence["limits"][0].update(vol_stock_max=str(Decimal(200)*FT3_TO_M3), vol_stock_remain=str(Decimal(100)*FT3_TO_M3))
    other = deepcopy(evidence["limits"][0])
    other.update(sid="3", sub_item_vol_unit_type=2, vol_stock_max="200", vol_stock_remain="100")
    evidence["limits"].append(other)
    source = service._context(dict(country="欧洲站", store="shop", msku="msku"), evidence)
    result = evaluate_capacity(DAY, 20, source)
    assert result["capacity_status"] == "source_available"
    assert result["capacity_remaining_m3"] == float(Decimal(100)*FT3_TO_M3)
    assert source["group_confirmed"] is False
    assert "未据此认定共享" in result["capacity_match_basis"]


def test_north_america_only_reads_own_us_site_not_other_countries():
    service = ReplenishmentCapacityService(Mock())
    evidence = sources()
    evidence["sellers"] += [dict(sid=4, name="shop-US", account_name="shop", seller_account_id=7, country="美国"),
                            dict(sid=5, name="shop-CA", account_name="shop", seller_account_id=7, country="加拿大")]
    quota = deepcopy(evidence["limits"][0])
    quota.update(sid="4", vol_stock_remain="2")
    evidence["limits"] += [quota, {**quota, "sid":"5", "vol_stock_remain":"9999"}]
    fee = deepcopy(evidence["fees"][0])
    fee.update(seller_name="shop-US", country="美国")
    evidence["fees"].append(fee)
    result = evaluate_capacity(DAY, 20, service._context(dict(country="北美站", store="shop", msku="msku"), evidence))
    assert result["capacity_remaining_m3"] == 2
    assert "美国" in result["capacity_quota_summary"] and "加拿大" not in result["capacity_quota_summary"]


def test_month_end_archive_that_is_not_month_end_is_not_reported_as_current_capacity():
    evidence = context()
    evidence.update(month="2026-09-01", source_time="2026-09-20", snapshot_time="2026-09-20",
                    archive_mode=True, reference_date=date(2026, 9, 30), demand_month="2026-09")
    result = evaluate_capacity(date(2026, 9, 10), 20, evidence)
    assert result["capacity_status"] == "pending"
    assert "未覆盖月底" in result["capacity_note"]


def test_ambiguous_store_accounts_cannot_borrow_quota():
    service = ReplenishmentCapacityService(Mock())
    evidence = sources()
    evidence["sellers"].append(dict(sid=6, name="another-FR", account_name="shop", seller_account_id=999, country="法国"))
    source = service._context(dict(country="欧洲站", store="shop", msku="msku"), evidence)
    assert source["fatal"] is True
    assert evaluate_capacity(DAY, 20, source)["capacity_remaining_m3"] is None


@pytest.mark.parametrize("store", ["Tboke", "QICE", "ouhao", "QINGLEE", "xinmu"])
def test_store_prefix_matches_existing_replenishment_etl_name_normalization(store):
    evidence = sources()
    evidence["groups"] = []
    evidence["sellers"][0].update(name=f"{store}-eu-DE", account_name=f"{store}-eu")
    evidence["fees"][0]["seller_name"] = f"{store}-eu-DE"
    source = ReplenishmentCapacityService(Mock())._context(dict(country="欧洲站", store=store, msku="msku"), evidence)
    assert source.get("fatal") is not True
    assert source["remaining"] == "-1"
    assert "sid=1" in source["match_basis"]


def test_prefix_match_still_rejects_ambiguous_accounts_within_the_same_category():
    evidence = sources()
    evidence["sellers"] += [dict(sid=6, name="name-eu-DE", account_name="name-eu", seller_account_id=66, country="德国"),
                            dict(sid=7, name="name-extra-FR", account_name="name-extra", seller_account_id=77, country="法国")]
    source = ReplenishmentCapacityService(Mock())._context(dict(country="欧洲站", store="name", msku="msku"), evidence)
    assert source["fatal"] is True
    assert "关联不唯一" in source["problems"][0]


def test_page_capacity_tooltip_only_contains_summary_one_action_and_date():
    evidence = context()
    evidence.update(remaining="25.721859", type_confirmed=False, type_label="大件（历史报表参考）",
                    match_basis="内部sid=2269，共享映射和重复占用核查过程", demand_complete=False)
    result = evaluate_capacity(DAY, 20, evidence)
    tooltip = result["capacity_tooltip"]
    assert len(tooltip.splitlines()) <= 4
    assert "25.72 m³" in tooltip and "大件" in tooltip and "2026-10-14" in tooltip
    assert "先核对商品仓储类型" in tooltip
    assert "sid=" not in tooltip and "同期新增" not in tooltip
    assert "内部sid=2269" in result["capacity_match_basis"]  # Full audit data survives.


def test_station_difference_tooltip_gives_a_single_clear_next_check():
    source = sources()
    source["groups"] = []
    source["sellers"].append(dict(sid=3, name="shop-FR", account_name="shop", seller_account_id=7, country="法国"))
    source["limits"].append({**source["limits"][0], "sid": "3", "vol_stock_remain": "3"})
    service = ReplenishmentCapacityService(Mock())
    result = evaluate_capacity(DAY, 20, service._context(dict(country="欧洲站", store="shop", msku="msku"), source))
    assert "先核对发往哪个国家" in result["capacity_tooltip"]
    assert "各站点额度数值" not in result["capacity_tooltip"]
    assert "德国" in result["capacity_quota_summary"] and "法国" in result["capacity_quota_summary"]


def test_source_failure_also_has_a_short_readable_tooltip():
    result = evaluate_capacity(DAY, 20, {"fatal": True, "problems": ["库容数据暂不可读取，请核对来源服务"]})
    assert "暂不可读取" in result["capacity_tooltip"]
    assert len(result["capacity_tooltip"].splitlines()) <= 4


def test_real_warehouse_type_overrides_size_tier_and_selects_its_own_quota():
    evidence = sources()
    evidence["warehouse_types"] = [dict(store_names="shop-DE,shop-FR", msku="msku", storage_type="服装")]
    evidence["limits"].append({**evidence["limits"][0], "sub_item_storage_type":"Apparel", "vol_stock_remain":"2"})
    source = ReplenishmentCapacityService(Mock())._context(dict(country="欧洲站", store="shop", msku="msku"), evidence)
    assert source["storage_type"] == "Apparel"
    assert source["type_confirmed"] is True
    assert source["remaining"] == "2"
    assert "FBA库存明细" in source["measurement_source"]


def test_new_plan_calculates_row_and_all_store_volume_without_changing_quantity():
    evidence = sources()
    evidence["groups"] = []
    evidence["warehouse_types"] = [dict(store_names="shop-DE", msku=key, storage_type="标准") for key in ("first","second")]
    evidence["limits"][0]["vol_stock_remain"] = ".03"
    evidence["fees"] = [{**evidence["fees"][0], "seller_sku":key} for key in ("first","second")]
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(return_value=evidence)
    universe = [dict(country="欧洲站", store="shop", msku=key, sku="sku", replenish_qty=20) for key in ("first","second")]
    row = deepcopy(universe[0])
    service.enrich([row], DAY, demand_rows=universe)
    assert row["replenish_qty"] == 20
    assert row["capacity_demand_volume_m3"] == .02
    assert row["capacity_group_demand_m3"] == .04
    assert row["capacity_shortfall_m3"] == .01
    assert row["capacity_status"] == "plan_group_over"
    assert "合计" in row["capacity_tooltip"] and "先核对商品" not in row["capacity_tooltip"]
    second_page = deepcopy(universe[1])
    service.enrich([second_page], DAY, demand_rows=universe)
    assert second_page["capacity_group_demand_m3"] == row["capacity_group_demand_m3"]


@pytest.mark.parametrize("qty,status", [(20,"plan_within"),(21,"plan_row_over"),(0,"no_replenishment")])
def test_new_plan_strict_boundary_and_zero_demand(qty,status):
    evidence = sources()
    evidence["warehouse_types"] = [dict(store_names="shop-DE", msku="msku", storage_type="标准")]
    evidence["limits"][0]["vol_stock_remain"] = ".02"
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(return_value=evidence)
    row = dict(country="欧洲站", store="shop", msku="msku", replenish_qty=qty)
    service.enrich([row], DAY, demand_rows=[deepcopy(row)])
    assert row["capacity_status"] == status
    assert row["replenish_qty"] == qty


def test_new_plan_uses_conservative_station_remaining_not_sum():
    evidence = sources()
    evidence["groups"] = []
    evidence["warehouse_types"] = [dict(store_names="shop-DE,shop-FR", msku="msku", storage_type="标准")]
    evidence["sellers"].append(dict(sid=3,name="shop-FR",account_name="shop",seller_account_id=7,country="法国"))
    evidence["limits"][0]["vol_stock_remain"] = ".03"
    evidence["limits"].append({**evidence["limits"][0],"sid":"3","vol_stock_remain":".02"})
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(return_value=evidence)
    row = dict(country="欧洲站",store="shop",msku="msku",replenish_qty=25)
    service.enrich([row], DAY, demand_rows=[deepcopy(row)])
    assert row["capacity_remaining_m3"] == .02
    assert row["capacity_status"] == "plan_row_over"
    assert "最小" in row["capacity_match_basis"]


def test_missing_type_or_volume_in_other_products_cannot_claim_store_plan_fits():
    evidence = sources()
    evidence["warehouse_types"] = [dict(store_names="shop-DE",msku="msku",storage_type="标准")]
    evidence["limits"][0]["vol_stock_remain"] = "2"
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(return_value=evidence)
    row = dict(country="欧洲站",store="shop",msku="msku",replenish_qty=20)
    unknown = dict(country="欧洲站",store="shop",msku="missing",replenish_qty=20)
    service.enrich([row], DAY, demand_rows=[deepcopy(row),unknown])
    assert row["capacity_status"] == "plan_incomplete"
    assert row["capacity_group_demand_m3"] == .02
    assert row["capacity_shortfall_m3"] is None


def test_same_msku_in_other_store_category_or_type_is_not_included_in_plan_group():
    evidence = sources()
    evidence["warehouse_types"] = [dict(store_names="shop-DE",msku="msku",storage_type="标准"),
                                   dict(store_names="shop-UK",msku="msku",storage_type="大件")]
    evidence["limits"] += [{**evidence["limits"][1],"sid":"2"}]
    evidence["limits"][0]["vol_stock_remain"] = "2"
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(return_value=evidence)
    row = dict(country="欧洲站",store="shop",msku="msku",replenish_qty=20)
    other = dict(country="英国站",store="shop",msku="msku",replenish_qty=100)
    service.enrich([row], DAY, demand_rows=[deepcopy(row),other])
    assert row["capacity_group_demand_m3"] == .02
    assert row["capacity_status"] == "plan_within"


def test_global_capacity_demand_uses_existing_display_quantity_and_no_page_role_filters():
    from app.services.replenishment_data import ReplenishmentDataService
    service = ReplenishmentDataService.__new__(ReplenishmentDataService)
    conn, cursor = MagicMock(), MagicMock()
    conn.cursor.return_value.__enter__.return_value = cursor
    cursor.fetchall.return_value = [dict(country_category="欧洲站",seller_name_new="shop",seller_sku_adj="msku",replenish_qty=50)]
    records = service._capacity_demand_items(conn, DAY)
    assert records[0]["replenish_qty"] == 50
    sql, params = cursor.execute.call_args.args
    assert "executable_replenish_qty" in sql and "max_cg_box_pcs" in sql
    assert "category" not in params and "keyword" not in params and "limit" not in sql
    assert params["snapshot_date"] == DAY


def test_listing_dimensions_support_mixed_edge_units_and_keep_small_volume():
    from app.services.replenishment_capacity import _listing_volume
    import json
    record = dict(dimension_info=json.dumps([dict(package_length=1,package_width=2,package_height=3,
                  package_length_units_type="centimeter",package_width_units_type="inches",package_height_units_type="centimeter")]))
    assert _listing_volume(record) == Decimal(".01")*Decimal(".0508")*Decimal(".03")


def test_plan_can_use_explicitly_labeled_fee_classification_estimate():
    evidence = sources()
    evidence["limits"][0]["vol_stock_remain"] = "2"
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(return_value=evidence)
    row = dict(country="欧洲站",store="shop",msku="msku",replenish_qty=20)
    service.enrich([row], DAY, demand_rows=[deepcopy(row)])
    assert row["capacity_status"] == "plan_within"
    assert "分类估算" in row["capacity_product_note"]
    assert "先核对商品" not in row["capacity_tooltip"]


def test_plan_with_dash_quota_keeps_volume_but_never_calls_placeholder_negative_overage():
    evidence = sources()
    evidence["limits"][0].update(vol_stock_max="-",vol_stock_remain="-1")
    evidence["warehouse_types"] = [dict(store_names="shop-DE",msku="msku",storage_type="标准")]
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(return_value=evidence)
    row = dict(country="欧洲站",store="shop",msku="msku",replenish_qty=20)
    service.enrich([row], DAY, demand_rows=[deepcopy(row)])
    assert row["capacity_status"] == "plan_unknown"
    assert row["capacity_demand_volume_m3"] == .02
    assert row["capacity_remaining_m3"] is None and row["capacity_shortfall_m3"] is None


def test_warehouse_type_conflict_is_not_silently_replaced_by_standard_fee_size():
    evidence = sources()
    evidence["warehouse_types"] = [dict(store_names="shop-DE",msku="msku",storage_type="标准"),
                                   dict(store_names="shop-DE",msku="msku",storage_type="服装")]
    context = ReplenishmentCapacityService(Mock())._context(dict(country="欧洲站",store="shop",msku="msku"),evidence)
    assert context["storage_type"] is None and context["type_conflict"] is True


def test_zero_new_demand_keeps_existing_negative_quota_warning_without_changing_quantity():
    evidence = sources()
    evidence["warehouse_types"] = [dict(store_names="shop-DE",msku="msku",storage_type="标准")]
    service = ReplenishmentCapacityService(Mock())
    service._load = Mock(return_value=evidence)
    row = dict(country="欧洲站",store="shop",msku="msku",replenish_qty=0)
    service.enrich([row], DAY, demand_rows=[])
    assert row["capacity_demand_volume_m3"] == 0 and row["replenish_qty"] == 0
    assert "已超占用" in row["capacity_status_label"]
    assert "剩余：-1" in row["capacity_tooltip"]
