"""Regression cases for the approved replenishment-only role rules."""
from datetime import date
from decimal import Decimal
from io import BytesIO
import sqlite3
import re

import pytest
from openpyxl import load_workbook

from app import main
from app.services.replenishment_data import ReplenishmentDataService
from app.services.replenishment_tracking_data import ReplenishmentTrackingService
from etl import replenishment_roles as roles
from etl import replenishment_update, replenishment_tracking_summary_update


def _unicode_ci(left, right):
    left, right = left.casefold(), right.casefold()
    return (left > right) - (left < right)


@pytest.mark.parametrize("daily,margin,want", [
    (0, None, "问题产品"), (0, .3, "问题产品"),
    (.5, .05, "瘦狗产品"), (.5, .049, "问题产品"),
    (1, .049, "问题产品"), (1.001, .049, "瘦狗产品"),
    (1, .05, "瘦狗产品"), (1, .10, "潜力产品"),
    (3, .25, "潜力产品"), (3, .251, "明星产品"),
    (5, .15, "潜力产品"), (5, .25, "潜力产品"),
    (5, .251, "明星产品"), (5.001, .15, "潜力产品"),
    (5.001, .151, "明星产品"), (6, .049, "瘦狗产品"),
    (2, None, None), (None, .3, None),
])
def test_role_sql_boundaries(daily, margin, want):
    with sqlite3.connect(":memory:") as conn:
        actual = conn.execute(
            f"select {roles.role_case('daily', 'margin')} from (select ? as daily, ? as margin)",
            (daily, margin),
        ).fetchone()[0]
    assert actual == want


def test_total_role_aggregates_own_countries_without_follow_compensation():
    with sqlite3.connect(":memory:") as conn:
        conn.create_collation("utf8mb4_unicode_ci", _unicode_ci)
        conn.execute("create table dashboard_pur_plan_replenish_data (cur_date text, country_category text, seller_name_new text, seller_sku_adj text, final_sales_30d real, daily_avg_sales real, asin_merge_flag int)")
        conn.execute("create table dashboard_replenishment_country_metrics (snapshot_date text, period_days int, country_category text, country text, seller_name_new text, seller_sku_adj text, sales_qty real, sales_amount real, order_gross_profit real)")
        conn.execute("insert into dashboard_pur_plan_replenish_data values ('2026-10-13','欧洲站','shop','sku',900,30,1)")
        conn.executemany("insert into dashboard_replenishment_country_metrics values (?,?,?,?,?,?,?,?,?)", [
            ('2026-10-13',30,'欧洲站','德国','shop','sku',45,1000,100),
            ('2026-10-13',30,'欧洲站','法国','shop','sku',15,3000,-20),
            ('2026-10-13',30,'北美站','美国','shop','sku',300,1000,500),
            ('2026-10-13',30,'欧洲站','德国','other','sku',300,1000,500),
            ('2026-10-12',30,'欧洲站','德国','shop','sku',300,1000,500),
            ('2026-10-13',7,'欧洲站','德国','shop','sku',100,1000,500),
        ])
        source = roles.replenishment_source_sql(30)
        metrics = ReplenishmentDataService.__new__(ReplenishmentDataService)._product_category_metric_sql(30)
        actual = conn.execute(f"select {metrics['daily_sales_expr']}, {metrics['margin_col']}, {metrics['category_expr']}, daily_avg_sales from {source} r").fetchone()
    assert actual == (2, .02, "瘦狗产品", 30)


def country(name, daily, margin, **overrides):
    result = dict(snapshot_date=date(2026,10,13), period_days=30,
                  country_category="欧洲站", seller_name_new="shop", seller_sku_adj="sku",
                  country=name, sales_qty=Decimal(str(daily))*30, sales_amount=Decimal("100"),
                  order_gross_profit=None if margin is None else Decimal(str(margin))*100)
    return dict(result, **overrides)


def row(**overrides):
    return dict(dict(cur_date=date(2026,10,13),country_category="欧洲站",
                     seller_name_new="shop",seller_sku_adj="sku",abcd_category="瘦狗产品",
                     role_daily_sales=Decimal("2"),role_margin_rate=Decimal(".02"),
                     replenish_qty=20), **overrides)


def test_review_lists_all_qualifying_countries_even_if_other_countries_missing():
    result = roles.low_margin_review(row(), [country("德国",2,.06),country("法国",3,.10)], 30)
    assert result["status"] == "qualified"
    assert "德国" in result["qualifying_country_metrics"]
    assert "法国" in result["qualifying_country_metrics"]
    assert "2.00" in result["qualifying_country_metrics"]
    assert "6.00%" in result["qualifying_country_metrics"]


def test_review_skips_countries_without_evaluable_sales_metrics():
    result = roles.low_margin_review(row(), [country("德国",2,.05),country("法国",2,None)],30)
    assert result["status"] == "none"
    assert "无国家站点同时满足：日均销量＞1件/天且毛利率＞5%" in result["low_margin_review_note"]
    assert "缺失" not in result["low_margin_review_note"]
    assert "数据不足" not in result["low_margin_review_note"]


def test_review_absent_country_means_not_operated_and_is_not_reported():
    result = roles.low_margin_review(row(), [country("德国",2,.04)],30)
    assert result["status"] == "none"
    assert "缺失" not in result["low_margin_review_note"]


def test_review_no_country_records_has_no_missing_country_warning():
    result = roles.low_margin_review(row(), [],30)
    assert result["status"] == "none"
    assert "数据不足" not in result["low_margin_review_note"]


def test_review_strict_boundaries_and_current_category_only():
    records = [country(name,1,.10) for name in ["德国","法国","意大利","西班牙","荷兰"]]
    records[0] = country("德国",2,.05)
    records += [country("美国",3,.20,country_category="北美站"),
                country("法国",3,.20,seller_name_new="other"),
                country("法国",3,.20,snapshot_date=date(2026,10,12)),
                country("法国",3,.20,period_days=7)]
    result = roles.low_margin_review(row(),records,30)
    assert result["status"] == "none"
    assert "无国家站点同时满足：日均销量＞1件/天且毛利率＞5%" in result["low_margin_review_note"]


@pytest.mark.parametrize("changes", [
    {"abcd_category":"明星产品"}, {"abcd_category":"潜力产品"},
    {"role_daily_sales":Decimal("1")}, {"role_margin_rate":Decimal(".05")},
    {"role_margin_rate":None},
])
def test_non_trigger_rows_are_unmarked(changes):
    result = roles.low_margin_review(row(**changes), [country("德国",2,.10)],30)
    assert result["status"] == ""
    assert result["low_margin_review_note"] == ""


@pytest.mark.parametrize("category,name", [("英国站","英国"),("北美站","美国")])
def test_single_country_categories(category,name):
    parent = row(country_category=category)
    result = roles.low_margin_review(parent,[country(name,1,.10,country_category=category)],30)
    assert result["status"] == "none"


def test_removed_180_day_period_falls_back_to_30():
    service = ReplenishmentDataService.__new__(ReplenishmentDataService)
    assert service._normalize_product_category_period_days(180) == 30
    assert service._normalize_product_category_period_days(90) == 90


def test_main_and_tracking_share_role_expression():
    service = ReplenishmentDataService.__new__(ReplenishmentDataService)
    tracking = ReplenishmentTrackingService.__new__(ReplenishmentTrackingService)
    for period in (7,14,30,90):
        expected = service._qualify_product_category_expr(service._product_category_metric_sql(period)["category_expr"],"p")
        assert tracking._product_category_expr(period) == expected


def test_export_review_font_colors_preserve_warning_fill_and_follow_border():
    columns = [{"name":name,"label":name} for name in ["seller_sku_adj","replenish_qty","low_margin_review_note","qualifying_country_metrics"]]
    rows = [dict(row(seller_sku_adj=status),support_replenish_level_sort=2,
                 final_sales_7d=20,final_sales_3d=20,fllow_flag=0,
                 low_margin_review_status=status,low_margin_review_note="核查说明",
                 qualifying_country_metrics="德国 日销2.00 毛利率6.00%")
            for status in ["qualified","none","insufficient"]]
    workbook = load_workbook(BytesIO(main.build_replenishment_xlsx({"columns":columns,"rows":rows})))
    sheet = workbook.active
    assert sheet.max_row == 4
    assert [sheet.cell(i,2).value for i in range(2,5)] == [20,20,20]
    for index,color in enumerate(["375623","9C0006","666666"],2):
        assert sheet.cell(index,1).fill.fgColor.rgb.endswith("FFF2CC")
        assert sheet.cell(index,3).fill.fgColor.rgb.endswith("FFF2CC")
        assert sheet.cell(index,4).fill.fgColor.rgb.endswith("FFF2CC")
        assert sheet.cell(index,3).font.color.rgb.endswith(color)
        assert sheet.cell(index,4).font.color.rgb.endswith(color)
        assert sheet.cell(index,sheet.max_column).border.right.style == "thick"


def test_export_legend_and_annotations_explain_overlapping_warnings_without_changing_values():
    columns = [{"name":name,"label":name} for name in ["seller_sku_adj","replenish_qty","low_margin_review_note","qualifying_country_metrics"]]
    source = dict(row(seller_sku_adj="OVERLAP"), seller_name_new="Disabled", support_replenish_level_sort=2,
                  final_sales_7d=20, final_sales_3d=20, fllow_flag=0,
                  low_margin_review_status="qualified", low_margin_review_note="总日均销量＞1件/天且总毛利率＜5%",
                  qualifying_country_metrics="德国 日销2.00 毛利率6.00%")
    payload = {"columns":columns,"rows":[source]}
    workbook = load_workbook(BytesIO(main.build_replenishment_xlsx(payload, {"disabled"})))
    assert workbook.sheetnames == ["补货计划", "颜色与标注说明"]
    sheet = workbook.active
    assert sheet.max_row == 2
    assert [sheet.cell(2,i).value for i in range(1,5)] == [source[c["name"]] for c in columns]
    assert len(columns) == 4  # Do not mutate the export payload.
    note = sheet.cell(2,5).value
    for marker in ["销量异常", "禁补店铺仍有补货量", "跟卖补货", "总维度低毛利", "日均销量＞1件/天且毛利率＞5%"]:
        assert marker in note
    assert all(cell.fill.fgColor.rgb.endswith("F4B183") for cell in sheet[2])
    assert sheet.cell(2,5).border.right.style == "thick"
    assert sheet.freeze_panes == "C2"
    assert sheet.auto_filter.ref == "A1:E2"
    legend = workbook.worksheets[-1]
    legend_text = "\n".join(str(cell.value or "") for cells in legend for cell in cells)
    for name in ["黄色底色", "红色底色", "橙色底色", "浅蓝底色", "深蓝粗边框", "绿色文字", "红色文字", "导出标注", "不调整补货数量"]:
        assert name in legend_text
    assert "达标" not in legend_text


def test_export_review_uses_text_without_new_fills_and_empty_exports_keep_legend():
    columns = [{"name":name,"label":name} for name in ["seller_sku_adj","replenish_qty","low_margin_review_note"]]
    rows = [dict(row(seller_sku_adj="REVIEW"), low_margin_review_status="none", low_margin_review_note="核查说明"),
            dict(row(seller_sku_adj="PLAIN"))]
    workbook = load_workbook(BytesIO(main.build_replenishment_xlsx({"columns":columns,"rows":rows})))
    sheet = workbook.active
    assert sheet.cell(2,3).fill.fill_type is None
    assert sheet.cell(2,3).font.color.rgb.endswith("9C0006")
    assert sheet.cell(3,4).value is None
    empty = load_workbook(BytesIO(main.build_replenishment_xlsx({"columns":[],"rows":[]})))
    assert empty.sheetnames == ["补货计划", "颜色与标注说明"]
    assert empty.active.max_row == 1


class SqliteConnection:
    """Execute emitted read queries, adapting only driver parameter syntax."""
    def __init__(self, conn):
        self.conn = conn
        self.queries = []

    def cursor(self):
        return self

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False

    def execute(self, query, params=None):
        self.queries.append(query)
        sql = re.sub(r"%\((\w+)\)s", r":\1", query)
        self.result = self.conn.execute(sql, params or {})

    def fetchall(self):
        return [dict(row) for row in self.result.fetchall()]


def test_export_service_keeps_selected_period_and_annotates_real_country_rows():
    with sqlite3.connect(":memory:") as db:
        db.create_collation("utf8mb4_unicode_ci", _unicode_ci)
        db.row_factory = sqlite3.Row
        db.execute("create table dashboard_pur_plan_replenish_data (cur_date text, country_category text, seller_name_new text, seller_sku_adj text, support_replenish_level_sort int, replenish_qty real)")
        db.execute("create table dashboard_replenishment_country_metrics (snapshot_date text, period_days int, country_category text, country text, seller_name_new text, seller_sku_adj text, sales_qty real, sales_amount real, order_gross_profit real)")
        db.execute("insert into dashboard_pur_plan_replenish_data values ('2026-10-13','欧洲站','shop','sku',2,20)")
        db.executemany("insert into dashboard_replenishment_country_metrics values (?,?,?,?,?,?,?,?,?)", [
            ('2026-10-13',7,'欧洲站','德国','shop','sku',14,100,10),
            ('2026-10-13',7,'欧洲站','法国','shop','sku',7,900,0),
            ('2026-10-13',30,'欧洲站','德国','shop','sku',60,100,40),
        ])
        conn = SqliteConnection(db)
        service = ReplenishmentDataService.__new__(ReplenishmentDataService)
        metrics = service._product_category_metric_sql(7)
        columns = ["cur_date","country_category","seller_name_new","seller_sku_adj","abcd_category","role_daily_sales","role_margin_rate"]
        rows = service._export_items(conn,"cur_date = %(snapshot_date)s",{"snapshot_date":"2026-10-13"},"category","asc",columns,metrics)
        service._attach_low_margin_reviews(conn,rows,7)
        assert len(rows) == 1
        assert rows[0]["role_daily_sales"] == 3
        assert rows[0]["abcd_category"] == "瘦狗产品"
        assert rows[0]["low_margin_review_status"] == "qualified"
        assert "德国" in rows[0]["qualifying_country_metrics"]
        assert "2.00" in rows[0]["qualifying_country_metrics"]


def test_persisted_30d_role_uses_the_same_new_case():
    sql = replenishment_update.REPLENISHMENT_RESULT_SQL
    # Execute the expression actually inserted by production ETL, not a second classifier.
    start = sql.index("    case\n        when (role_sales_30 / 30.0)")
    expr = sql[start:sql.index(" as abcd_category",start)].strip()
    with sqlite3.connect(":memory:") as db:
        assert db.execute(f"select {expr} from (select 60 as role_sales_30, 0.04 as pprofit_ratio_30)").fetchone()[0] == "瘦狗产品"


def test_summary_etl_classifies_own_metrics_instead_of_stored_old_role():
    with sqlite3.connect(":memory:") as db:
        expr = replenishment_tracking_summary_update.PRODUCT_CATEGORY_30D_SQL
        value = db.execute(f"select {expr} from (select 2 as role_daily_sales_30d, 0.04 as role_margin_rate_30d, '问题产品' as abcd_category) l").fetchone()[0]
    assert value == "瘦狗产品"


def test_etl_missing_own_sales_evidence_does_not_become_zero_sales():
    sql = replenishment_update.REPLENISHMENT_RESULT_SQL
    end = sql.index(" as abcd_category")
    expr = sql[sql.rfind("    case", 0, end):end].strip()
    with sqlite3.connect(":memory:") as db:
        result = db.execute(f"select {expr} from (select 0 as sales_30, null as role_sales_30, null as pprofit_ratio_30)").fetchone()[0]
    assert result is None


def test_tracking_summary_only_aggregates_snapshots_in_current_cutoff():
    reads = []
    with sqlite3.connect(":memory:") as db:
        db.create_collation("utf8mb4_unicode_ci", _unicode_ci)
        db.create_collation("utf8mb4_0900_ai_ci", _unicode_ci)
        db.row_factory = sqlite3.Row
        db.create_function("observed_sales",1,lambda value: reads.append(value) or value)
        db.execute("create table country_source (snapshot_date text, period_days int, country_category text, seller_name_new text, seller_sku_adj text, sales_qty real, sales_amount real, order_gross_profit real)")
        db.execute("create view dashboard_replenishment_country_metrics as select snapshot_date, period_days, country_category, seller_name_new, seller_sku_adj, observed_sales(sales_qty) as sales_qty, sales_amount, order_gross_profit from country_source")
        db.execute("create table dashboard_replenishment_tracking_summary (cutoff_date text, latest_replenishment_date text, country_category text, seller_name_new text, seller_sku_adj text)")
        db.execute("insert into dashboard_replenishment_tracking_summary values ('2026-10-13','2026-10-12','欧洲站','shop','sku')")
        db.executemany("insert into country_source values (?,30,'欧洲站','shop','sku',60,100,4)", [(day,) for day in ['2026-10-12','2026-10-11','2026-10-10']])
        db.executemany(
            "insert into country_source values ('2026-10-12',30,'欧洲站','other',?,60,100,4)",
            [(f"unrelated-{index}",) for index in range(5)],
        )
        conn = SqliteConnection(db)
        conn.execute(f"select current_role_category from {roles.tracking_summary_source_sql(30)} s", {"cutoff_date":"2026-10-13"})
        assert conn.fetchall()[0]["current_role_category"] == "瘦狗产品"
    assert len(reads) == 1
