import unittest
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import main


class FakeDetailExportService:
    def get_detail_export_payload(self, filters):
        return {
            "columns": ["snapshot_date", "seller_name_new", "sales_qty"],
            "rows": [
                {
                    "snapshot_date": "2026-06-15",
                    "seller_name_new": "booyee",
                    "sales_qty": 12,
                }
            ],
            "start_date": "2026-06-01",
            "end_date": "2026-06-15",
        }


class FakeReplenishmentExportService:
    def get_non_replenishment_stores(self):
        return set()

    def get_export_payload(self, **_filters):
        columns = [
            {"name": "fba_local_quantity", "label": "FBA本地库存"},
            {"name": "total", "label": "总库存"},
            {"name": "available_total", "label": "可用库存"},
            {"name": "afn_fulfillable_quantity", "label": "FBA可售库存"},
            {"name": "stock_up_num", "label": "在途数量"},
            {"name": "afn_unsellable_quantity", "label": "FBA不可售库存"},
            {"name": "sc_quantity_local_valid", "label": "本地可用库存"},
            {"name": "sc_quantity_purchase_shipping", "label": "采购在途数量"},
            {"name": "sc_quantity_purchase_plan", "label": "采购计划数量"},
            {"name": "sc_quantity_local_qc", "label": "本地质检数量"},
            {"name": "local_quantity", "label": "本地库存"},
        ]
        return {
            "columns": columns,
            "rows": [
                {
                    column["name"]: Decimal(f"{index}.0000")
                    for index, column in enumerate(columns, start=1)
                }
            ],
            "snapshot_date": "2026-07-15",
        }


class ExportHeaderTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(main.app)
        self.service = FakeDetailExportService()

    def test_detail_export_uses_chinese_headers(self):
        with patch("app.main.dashboard_service", self.service):
            response = self.client.get("/api/detail/export")

        self.assertEqual(response.status_code, 200)
        first_line = response.text.splitlines()[0].lstrip("\ufeff")
        self.assertEqual(first_line, "快照日期,店铺,销量")

    def test_replenishment_export_writes_business_definition_comments_to_xlsx_headers(self):
        expected_comments = {
            "FBA本地库存": "FBA本地 = FBA总库存+本地库存",
            "总库存": "FBA总库存 = FBA可售+FBA预留+待调仓+标发在途+入库中",
            "可用库存": "FBA可用库存 = FBA可售+待调仓+FBA预留+入库中（此字段可在业务配置中自定义）",
            "FBA可售库存": "FBA可售库存 = afn fulfillable",
            "在途数量": "FBA在途 = 实际在途发货单发货数量-签收数量",
            "FBA不可售库存": "FBA不可售 = Unfulfillable",
            "本地可用库存": "本地可用 = 配对SKU的可用量+可用锁定量+期望可用量，仅组合产品包含期望可用量；点击设置-业务配置-补货建议中自定义",
            "采购在途数量": "采购在途 = 目的仓为本地仓的调拨单待收货量",
            "采购计划数量": "采购计划 = 配对SKU相关采购计划单待采购量统计数据：采购计划（待审批）+采购计划（待采购）",
            "本地质检数量": "本地质检 = 配对SKU的待检待上架量（汇总SKU无绑定FNSKU数量与SKU+FNSKU数量）",
            "本地库存": "本地库存 = 本地可用+采购在途+采购计划+本地质检",
        }

        with patch("app.main.replenishment_service", FakeReplenishmentExportService()):
            response = self.client.get("/api/replenishment/export")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("replenishment_2026-07-15.xlsx", response.headers["content-disposition"])

        workbook = load_workbook(BytesIO(response.content), read_only=False)
        worksheet = workbook.active
        for cell in worksheet[1]:
            self.assertIsNotNone(cell.comment, cell.value)
            self.assertEqual(cell.comment.text, expected_comments[cell.value])
        self.assertEqual([cell.value for cell in worksheet[2]], list(range(1, 12)))
        self.assertTrue(all(cell.data_type == "n" for cell in worksheet[2]))

    def test_replenishment_sales_concentration_rule(self):
        cases = [
            ({"final_sales_3d": 7, "final_sales_7d": 10, "support_replenish_level_sort": 1, "replenish_qty": 50}, True),
            (
                {
                    "final_sales_3d": 7,
                    "final_sales_7d": 10,
                    "support_replenish_level_sort": 1,
                    "asin_merge_flag": "否",
                    "replenish_qty": 50,
                },
                True,
            ),
            ({"final_sales_3d": Decimal("6.9"), "final_sales_7d": 10, "support_replenish_level_sort": 2}, False),
            ({"final_sales_3d": 7, "final_sales_7d": 9, "support_replenish_level_sort": 3}, False),
            ({"final_sales_3d": 0, "final_sales_7d": 0, "support_replenish_level_sort": 1}, False),
            ({"final_sales_3d": 70, "final_sales_7d": 100, "support_replenish_level_sort": 4}, False),
            (
                {
                    "final_sales_3d": 7,
                    "final_sales_7d": 10,
                    "support_replenish_level_sort": 1,
                    "asin_merge_flag": 1,
                    "replenish_qty": 0,
                },
                False,
            ),
        ]

        for row, expected in cases:
            with self.subTest(row=row):
                self.assertEqual(expected, main.is_replenishment_sales_concentrated(row))

    def test_mad_spike_highlight_requires_active_level_and_positive_replenishment(self):
        cases = [
            ({"support_replenish_level_sort": 1, "replenish_qty": 2295, "sales_spike_flag": 1, "sales_spike_status": "suspected"}, True),
            ({"support_replenish_level_sort": 2, "replenish_qty": 2295, "sales_spike_flag": 1, "sales_spike_status": "confirmed_recovered"}, True),
            ({"support_replenish_level_sort": 3, "replenish_qty": 0, "sales_spike_flag": 1}, False),
            ({"support_replenish_level_sort": 4, "replenish_qty": 2295, "sales_spike_flag": 1}, False),
            ({"support_replenish_level_sort": 1, "replenish_qty": 2295, "sales_spike_flag": 0, "sales_spike_status": "sustained_growth", "final_sales_3d": 5, "final_sales_7d": 183}, False),
        ]

        for row, expected in cases:
            with self.subTest(row=row):
                self.assertEqual(expected, main.is_replenishment_sales_concentrated(row))

    def test_mad_spike_xlsx_keeps_quantity_and_hides_metadata_columns(self):
        payload = {
            "columns": [
                {"name": "seller_sku_adj", "label": "MSKU"},
                {"name": "replenish_qty", "label": "补货数量"},
            ],
            "rows": [
                {
                    "seller_sku_adj": "KQ-UFY7-HWK2",
                    "replenish_qty": Decimal("2295"),
                    "support_replenish_level_sort": 1,
                    "sales_spike_flag": 1,
                    "sales_spike_status": "confirmed_recovered",
                }
            ],
        }

        workbook = load_workbook(BytesIO(main.build_replenishment_xlsx(payload)), read_only=False)
        worksheet = workbook.active

        self.assertEqual(["MSKU", "补货数量"], [cell.value for cell in worksheet[1]])
        self.assertEqual(2295, worksheet["B2"].value)
        self.assertTrue(all(cell.fill.fgColor.rgb.endswith("FFF2CC") for cell in worksheet[2]))

    def test_replenishment_xlsx_hides_secondary_export_columns_without_removing_values(self):
        payload = {
            "columns": [
                {"name": "seller_sku_adj", "label": "MSKU"},
                {"name": "replenish_qty", "label": "补货数量"},
                {"name": "fba_local_quantity", "label": "FBA本地库存"},
                {"name": "stock_up_num", "label": "在途数量"},
                {"name": "purchase_lead_status", "label": "采购交期状态"},
                {"name": "supplier_moq", "label": "供应商最小起订量"},
            ],
            "rows": [
                {
                    "seller_sku_adj": "MSKU-1",
                    "replenish_qty": 30,
                    "fba_local_quantity": 100,
                    "stock_up_num": 20,
                    "purchase_lead_status": "正常",
                    "supplier_moq": 50,
                }
            ],
        }

        workbook = load_workbook(BytesIO(main.build_replenishment_xlsx(payload)), read_only=False)
        worksheet = workbook.active

        self.assertEqual(100, worksheet["C2"].value)
        self.assertEqual(20, worksheet["D2"].value)
        self.assertTrue(worksheet.column_dimensions["C"].hidden)
        self.assertTrue(worksheet.column_dimensions["D"].hidden)
        self.assertTrue(worksheet.column_dimensions["E"].hidden)
        self.assertTrue(worksheet.column_dimensions["F"].hidden)
        self.assertFalse(worksheet.column_dimensions["A"].hidden)
        self.assertFalse(worksheet.column_dimensions["B"].hidden)

    def test_replenishment_xlsx_hides_long_cycle_and_audit_columns(self):
        hidden_column_names = [
            "onsale_sites",
            "unsale_sites",
            "marketplace_concat",
            "predict_abcd_category",
            "r_90d_salable_days",
            "sales_180d",
            "sales_90d",
            "amount_180d",
            "amount_90d",
            "pprofit_180d",
            "pprofit_90d",
            "pprofit_ratio_180d",
            "pprofit_ratio_90d",
            "60d_stocko_qty",
            "90d_stocko_qty",
            "180d_stocko_qty",
            "amz_instock_sales_ratio",
            "instock_intrans_pur_sales_ratio",
            "created_at",
            "updated_at",
        ]
        payload = {
            "columns": [
                {"name": "seller_sku_adj", "label": "MSKU"},
                *[{"name": name, "label": name} for name in hidden_column_names],
            ],
            "rows": [{"seller_sku_adj": "MSKU-1", **{name: 1 for name in hidden_column_names}}],
        }

        workbook = load_workbook(BytesIO(main.build_replenishment_xlsx(payload)), read_only=False)
        worksheet = workbook.active

        for column_index in range(2, len(payload["columns"]) + 1):
            self.assertEqual(1, worksheet.cell(2, column_index).value)
            self.assertTrue(worksheet.column_dimensions[worksheet.cell(1, column_index).column_letter].hidden)
        self.assertFalse(worksheet.column_dimensions["A"].hidden)

    def test_replenishment_xlsx_highlights_concentrated_sales_row(self):
        payload = {
            "columns": [
                {"name": "seller_sku_adj", "label": "MSKU"},
                {"name": "final_sales_3d", "label": "3天销量"},
                {"name": "final_sales_7d", "label": "7天销量"},
            ],
            "rows": [
                {
                    "seller_sku_adj": "SPIKE-1",
                    "final_sales_3d": Decimal("7"),
                    "final_sales_7d": Decimal("10"),
                    "support_replenish_level_sort": 1,
                    "asin_merge_flag": 0,
                    "replenish_qty": 50,
                },
                {
                    "seller_sku_adj": "NORMAL-1",
                    "final_sales_3d": Decimal("6"),
                    "final_sales_7d": Decimal("10"),
                    "support_replenish_level_sort": 1,
                    "asin_merge_flag": 0,
                    "replenish_qty": 50,
                },
            ],
        }

        workbook = load_workbook(BytesIO(main.build_replenishment_xlsx(payload)), read_only=False)
        worksheet = workbook.active

        self.assertIsNotNone(worksheet["B1"].comment)
        self.assertIn("最新完整销售日疑似单日爆单", worksheet["B1"].comment.text)
        self.assertIn("不修改补货数量", worksheet["B1"].comment.text)
        self.assertTrue(all(cell.fill.fill_type == "solid" for cell in worksheet[2]))
        self.assertTrue(all(cell.fill.fgColor.rgb.endswith("FFF2CC") for cell in worksheet[2]))
        self.assertTrue(all(cell.fill.fill_type is None for cell in worksheet[3]))

    def test_disabled_replenishment_store_requires_positive_exported_quantity(self):
        disabled_stores = {"tboke"}
        cases = [
            ({"seller_name_new": "Tboke", "replenish_qty": 50}, True),
            ({"seller_name_new": "TBOKE", "replenish_qty": Decimal("0.01")}, True),
            ({"seller_name_new": "Tboke", "replenish_qty": 0}, False),
            ({"seller_name_new": "booyee", "replenish_qty": 50}, False),
            ({"seller_name_new": "", "replenish_qty": 50}, False),
        ]

        for row, expected in cases:
            with self.subTest(row=row):
                self.assertEqual(
                    expected,
                    main.is_disabled_replenishment_store_with_qty(row, disabled_stores),
                )

    def test_replenishment_xlsx_uses_red_and_orange_without_changing_yellow_rule(self):
        payload = {
            "columns": [
                {"name": "seller_name_new", "label": "Store"},
                {"name": "seller_sku_adj", "label": "MSKU"},
                {"name": "replenish_qty", "label": "Replenish Qty"},
            ],
            "rows": [
                {
                    "seller_name_new": "Tboke",
                    "seller_sku_adj": "RED-1",
                    "replenish_qty": 50,
                    "support_replenish_level_sort": 1,
                    "final_sales_3d": 1,
                    "final_sales_7d": 10,
                },
                {
                    "seller_name_new": "Tboke",
                    "seller_sku_adj": "ORANGE-1",
                    "replenish_qty": 50,
                    "support_replenish_level_sort": 1,
                    "final_sales_3d": 7,
                    "final_sales_7d": 10,
                },
                {
                    "seller_name_new": "booyee",
                    "seller_sku_adj": "YELLOW-1",
                    "replenish_qty": 50,
                    "support_replenish_level_sort": 1,
                    "final_sales_3d": 7,
                    "final_sales_7d": 10,
                },
                {
                    "seller_name_new": "Tboke",
                    "seller_sku_adj": "NORMAL-1",
                    "replenish_qty": 0,
                    "support_replenish_level_sort": 1,
                    "final_sales_3d": 1,
                    "final_sales_7d": 10,
                },
            ],
        }

        workbook = load_workbook(
            BytesIO(main.build_replenishment_xlsx(payload, {"tboke"})),
            read_only=False,
        )
        worksheet = workbook.active

        self.assertTrue(all(cell.fill.fgColor.rgb.endswith("F4CCCC") for cell in worksheet[2]))
        self.assertTrue(all(cell.fill.fgColor.rgb.endswith("F4B183") for cell in worksheet[3]))
        self.assertTrue(all(cell.fill.fgColor.rgb.endswith("FFF2CC") for cell in worksheet[4]))
        self.assertTrue(all(cell.fill.fill_type is None for cell in worksheet[5]))


if __name__ == "__main__":
    unittest.main()
