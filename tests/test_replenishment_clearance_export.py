"""Keep clearance evidence small, separate, dated and explicit about old snapshots."""
from datetime import date
from io import BytesIO
import re
import sqlite3
from unittest.mock import MagicMock

import pymysql
from openpyxl import load_workbook
import pytest

from app import main
from app.services.replenishment_data import ReplenishmentDataService
from etl import replenishment_update as update


def workbook(audit):
    return load_workbook(BytesIO(main.build_replenishment_xlsx({
        'columns': [{'name': 'seller_sku_adj', 'label': 'MSKU'},
                    {'name': 'replenish_qty', 'label': '补货数量'}],
        'rows': [{'seller_sku_adj': 'normal', 'replenish_qty': 100}],
        'clearance_rows': audit, 'snapshot_date': '2026-10-13',
    })))


def test_nonempty_audit_is_only_five_columns_and_never_expands_the_main_sheet():
    wb = workbook([{'country_category': '欧洲站', 'seller_name_new': 'Shop',
                     'seller_sku_adj': 'blocked', 'reference_qty': 50,
                     'block_reason': '清货名单，禁止补货'}])
    assert wb.sheetnames == ['补货计划', '清货拦截明细', '颜色与标注说明']
    assert wb['补货计划'].max_row == 2
    assert wb['补货计划'].max_column == 3
    assert wb['补货计划']['B2'].value == 100
    sheet = wb['清货拦截明细']
    assert sheet.max_column == 5
    assert sheet['C2'].value == 'blocked'
    assert sheet['D2'].value == 50
    assert sheet['E2'].value == '清货名单，禁止补货'
    assert '参考' in sheet['D1'].value


@pytest.mark.parametrize('rows', [[], [{'reference_qty': 0}], [{'reference_qty': -1}]])
def test_no_positive_clearance_does_not_create_an_empty_sheet(rows):
    assert workbook(rows).sheetnames == ['补货计划', '颜色与标注说明']


def mock_conn(rows):
    conn = MagicMock()
    conn.__enter__.return_value = conn
    cursor = conn.cursor.return_value.__enter__.return_value
    cursor.fetchall.return_value = rows
    return conn, cursor


def test_saved_evidence_is_preserved_when_current_list_is_empty():
    service = ReplenishmentDataService.__new__(ReplenishmentDataService)
    source, _ = mock_conn([])
    service.connect_source = lambda: source
    conn, cursor = mock_conn([{'reference_qty': 25}])
    assert service._clearance_export_items(conn, date(2026, 10, 15), '欧洲站', 'Shop', 'ABC') == [{'reference_qty': 25}]
    sql, params = cursor.execute.call_args.args
    assert 'snapshot_date' in sql and 'reference_qty > 0' in sql
    assert params['snapshot_date'] == date(2026, 10, 15)
    assert params['site'] == '欧洲站' and params['store'] == 'Shop'
    assert params['keyword'] == '%ABC%'
    assert 'role_' not in sql


def test_saved_evidence_does_not_skip_new_clearance_entries_on_same_snapshot():
    service = ReplenishmentDataService.__new__(ReplenishmentDataService)
    saved = {'country_category': '欧洲站', 'seller_name_new': 'Shop',
             'seller_sku_adj': 'old', 'reference_qty': 25}
    added = {'country_category': '欧洲站', 'seller_name_new': 'Shop',
             'seller_sku_adj': 'new', 'reference_qty': 36}
    conn, cursor = mock_conn([])
    cursor.fetchall.side_effect = [[saved], [dict(added)]]
    source, _ = mock_conn([{'msku': 'new', 'store': '全部'}])
    service.connect_source = lambda: source
    rows = service._clearance_export_items(conn, date(2026, 10, 15), 'all', 'all', '')
    assert [(row['seller_sku_adj'], row['reference_qty']) for row in rows] == [
        ('old', 25), ('new', 36)]


def test_old_snapshot_clearance_catches_calculated_quantity_below_moq():
    db = sqlite3.connect(':memory:')
    db.row_factory = sqlite3.Row
    db.execute('create table dashboard_replenishment_clearance_audit '
               '(snapshot_date text, country_category text, seller_name_new text, '
               'seller_sku_adj text, reference_qty real)')
    db.execute('create table dashboard_pur_plan_replenish_data '
               '(cur_date text, country_category text, seller_name_new text, seller_sku_adj text, '
               'moq_status text, calculated_replenish_qty real, executable_replenish_qty real, '
               'history_recovery_flag int, support_replenish_level_sort int, '
               'replenish_block_reason text, asin_merge_flag int, replenish_qty real, max_cg_box_pcs real)')
    db.execute("insert into dashboard_pur_plan_replenish_data values "
               "('2026-10-15','欧洲站','Shop','ABC','below_minimum',36,0,0,1,'',0,0,12)")
    conn, cursor = mock_conn([])
    results = []

    def execute(sql, params):
        sqlite_params = {key: value.isoformat() if isinstance(value, date) else value
                         for key, value in params.items()}
        results[:] = [dict(row) for row in db.execute(re.sub(r'%\((\w+)\)s', r':\1', sql), sqlite_params)]

    cursor.execute.side_effect = execute
    cursor.fetchall.side_effect = lambda: list(results)
    source, _ = mock_conn([{'msku': 'ABC', 'store': '全部'}])
    service = ReplenishmentDataService.__new__(ReplenishmentDataService)
    service.connect_source = lambda: source
    try:
        rows = service._clearance_export_items(conn, date(2026, 10, 15), 'all', 'all', '')
        assert [(row['seller_sku_adj'], row['reference_qty']) for row in rows] == [('ABC', 36)]
    finally:
        db.close()


def test_old_snapshot_evidence_uses_original_display_quantity():
    service = ReplenishmentDataService.__new__(ReplenishmentDataService)
    conn, cursor = mock_conn([{'country_category': '欧洲站', 'seller_name_new': 'Shop',
                             'seller_sku_adj': 'ABC', 'reference_qty': 36}])
    cursor.execute.side_effect = [pymysql.err.ProgrammingError(1146, 'audit table absent'), None]
    source, source_cursor = mock_conn([{'msku': 'ABC', 'store': '全部'},
                                       {'msku': 'DEF', 'store': 'shop'}])
    service.connect_source = MagicMock(return_value=source)
    rows = service._clearance_export_items(conn, date(2026, 10, 13), 'all', 'all', '')
    assert rows[0]['reference_qty'] == 36
    assert rows[0]['block_reason'] == '清货名单，禁止补货'
    sql, params = cursor.execute.call_args.args
    assert 'executable_replenish_qty' in sql
    assert 'history_recovery_flag' in sql
    assert params['clearance_msku_0'] == 'ABC'
    assert params['clearance_msku_1'] == 'DEF'
    assert params['clearance_store_1'] == 'shop'
    assert 'r.cur_date=' in sql
    assert 'reference_qty > 0' in sql


@pytest.mark.parametrize('filters', [
    {}, {'keyword': 'Shop'}, {'keyword': 'SKU-A'}, {'keyword': '欧洲站'},
    {'site': '欧洲站', 'store': 'Shop', 'keyword': 'ABC'},
])
def test_old_snapshot_export_zeroes_blocked_quantities_before_capacity_and_workbook(monkeypatch, filters):
    service = ReplenishmentDataService.__new__(ReplenishmentDataService)
    conn, _ = mock_conn([])
    service.connect = lambda: conn
    service._export_columns = lambda _: [
        {'name': name, 'label': name} for name in (
            'seller_name_new', 'seller_sku_adj', 'replenish_qty',
            'replenish_box_qty', 'replenish_cost', 'replenish_block_reason')]
    fields = ('replenish_qty', 'replenish_box_qty', 'replenish_cost',
              'calculated_replenish_qty', 'calculated_replenish_box_qty',
              'calculated_replenish_cost', 'executable_replenish_qty',
              'executable_replenish_box_qty', 'executable_replenish_cost',
              'replenish_need_qty', 'base_replenish_need_qty',
              'lead_adjusted_replenish_need_qty', 'history_recovery_need_qty',
              'moq_shortfall_qty')
    blocked = dict.fromkeys(fields, 50)
    blocked.update(country_category='欧洲站', seller_name_new='Shop', seller_sku_adj='ABC',
                   history_recovery_flag=1, moq_status='已达起订量', available_total=7)
    rows = [
        blocked,
        {**blocked, 'seller_name_new': 'Other'},
        {**blocked, 'country_category': '英国站'},
        {**blocked, 'seller_sku_adj': 'normal'},
    ]
    service._export_items = lambda *a, **kw: [dict(row) for row in rows]
    service._attach_low_margin_reviews = lambda *a: None
    service._attach_product_tags = lambda *a: None
    outside = {**blocked, 'seller_name_new': 'Outside', 'seller_sku_adj': 'OUT'}
    service._capacity_demand_items = lambda *a: [dict(row) for row in [*rows, outside]]
    all_clearance = [
        {'country_category': '欧洲站', 'seller_name_new': 'Shop', 'seller_sku_adj': 'ABC',
         'reference_qty': 50, 'block_reason': '清货名单，禁止补货'},
        {'country_category': '欧洲站', 'seller_name_new': 'Outside', 'seller_sku_adj': 'OUT',
         'reference_qty': 50, 'block_reason': '清货名单，禁止补货'}]

    def clearance(_conn, _date, site, store, keyword):
        if site == store == 'all' and keyword == '':
            return all_clearance
        return all_clearance[:1] if keyword in ('', 'ABC') else []

    service._clearance_export_items = clearance
    capacity_inputs = []

    def enrich(_self, items, selected_date, *, demand_rows):
        capacity_inputs.append((items, demand_rows))

    monkeypatch.setattr('app.services.replenishment_data.ReplenishmentCapacityService.enrich', enrich)
    payload = service.get_export_payload(snapshot_date='2026-10-13', **filters)
    row = payload['rows'][0]
    assert {field: row[field] for field in fields} == dict.fromkeys(fields, 0)
    assert row['history_recovery_flag'] == 0
    assert row['moq_status'] == '补货数量为0'
    assert row['replenish_block_reason'] == '清货名单，禁止补货'
    assert row['available_total'] == 7
    assert [r['replenish_qty'] for r in payload['rows'][1:]] == [50, 50, 50]
    assert [r['replenish_qty'] for r in capacity_inputs[0][1]] == [0, 50, 50, 50, 0]
    wb = load_workbook(BytesIO(main.build_replenishment_xlsx(payload)))
    assert [wb['补货计划'].cell(2, col).value for col in (3, 4, 5)] == [0, 0, 0]
    assert '清货' in wb['补货计划'].cell(2, wb['补货计划'].max_column).value
    if payload['clearance_rows']:
        assert wb['清货拦截明细']['D2'].value == 50


def test_errors_other_than_missing_audit_table_are_not_silenced():
    service = ReplenishmentDataService.__new__(ReplenishmentDataService)
    conn, cursor = mock_conn([])
    cursor.execute.side_effect = pymysql.err.OperationalError(2013, 'connection lost')
    with pytest.raises(pymysql.err.OperationalError):
        service._clearance_export_items(conn, date(2026, 10, 15), 'all', 'all', '')


def test_evidence_uses_original_quantity_with_history_recovery_moq_and_store_scope():
    db = sqlite3.connect(':memory:')
    db.create_function('greatest', 2, max)
    db.execute("attach database ':memory:' as etl_datasync_test")
    db.execute('create table clearance_reference_original (country_category text, seller_name_new text, seller_sku_adj text, max_sku text, history_recovery_flag int, support_replenish_level_sort int, replenish_qty real, max_cg_box_pcs real, cur_date text default "2026-10-15", replenish_block_reason text default "", asin_merge_flag int default 0)')
    db.execute('create table etl_datasync_test.dashboard_replenishment_clearance_sync (msku text, store text)')
    db.execute('create table etl_datasync_test.dashboard_replenishment_supplier_moq_sync (snapshot_date text, sku text, supplier_moq real)')
    db.execute('create table etl_datasync_test.dashboard_replenishment_clearance_audit (snapshot_date text, country_category text, seller_name_new text, seller_sku_adj text, reference_qty real)')
    db.executemany('insert into clearance_reference_original (country_category,seller_name_new,seller_sku_adj,max_sku,history_recovery_flag,support_replenish_level_sort,replenish_qty,max_cg_box_pcs) values (?,?,?,?,?,?,?,?)', [
        ('欧洲站', 'Shop', 'normal', 'A', 0, 1, 80, 10),
        ('欧洲站', 'Shop', 'clear', 'A', 0, 1, 80, 10),
        ('英国站', 'Shop', 'clear', 'B', 0, 1, 20, 10),
        ('欧洲站', 'Shop', 'recover', 'C', 1, 4, 0, 36),
        ('欧洲站', 'Shop', 'zero', 'D', 0, 4, 0, 10),
        ('欧洲站', 'Shop', 'followed', 'D', 0, 1, 0, 10),
        ('欧洲站', 'Other', 'specific', 'D', 0, 1, 80, 10),
        ('欧洲站', 'Shop', 'specific', 'D', 0, 1, 80, 10),
    ])
    db.executemany('insert into etl_datasync_test.dashboard_replenishment_clearance_sync values (?,?)',
                   [('clear', '全部'), ('recover', '全部'), ('zero', '全部'),
                    ('followed', '全部'), ('specific', 'shop')])
    db.execute("insert into etl_datasync_test.dashboard_replenishment_supplier_moq_sync values ('2026-10-15','B',50)")
    # Execute the real production quantity/filter projection over original-result
    # fixtures. ASIN target restoration is executed in the follow-store tests.
    projection = update.build_clearance_audit_sql().split('clearance_reference_quantities as (', 1)[1]
    sql = ('insert into etl_datasync_test.dashboard_replenishment_clearance_audit '
           'with clearance_reference_quantities as (' + projection)
    sql = sql.replace('%(snapshot_date)s', ':snapshot_date').replace('binary ', '')
    db.execute(sql, {'snapshot_date': '2026-10-15'})
    rows = db.execute('select seller_sku_adj, reference_qty from etl_datasync_test.dashboard_replenishment_clearance_audit order by seller_sku_adj').fetchall()
    assert rows == [('clear', 80), ('recover', 36), ('specific', 80)]
    db.close()


def test_audit_reuses_existing_candidates_and_replaces_only_one_snapshot_atomically():
    schemas = update.SchemaConfig(
        target_schema='etl_datasync_test', etl_source_schema='etl_datasync',
        dwd_source_schema='dwd_datasync', pricing_source_schema='temporary_dwd')
    conn, cursor = mock_conn([])
    update.refresh_clearance_audit(conn, schemas, {'snapshot_date': date(2026, 10, 15)})
    delete, insert = [call.args[0] for call in cursor.execute.call_args_list]
    assert 'where snapshot_date = %(snapshot_date)s' in delete
    assert 'dashboard_replenishment_work_replenish_calc_v3' in insert
    assert 'tmp_pur_plan_replenish_calc' not in insert
    conn.commit.assert_called_once()
    assert 'dashboard_replenishment_clearance_audit' in '\n'.join(update.DDL_STATEMENTS)


def test_result_must_validate_before_exclusion_evidence_is_written(monkeypatch):
    conn, cursor = mock_conn([])
    cursor.rowcount = 0
    events = []
    monkeypatch.setattr(update, 'build_replenishment_result_statements', lambda _: [])
    monkeypatch.setattr(update, 'refresh_replenishment_sales_spike_flags', lambda *a, **kw: 0)
    monkeypatch.setattr(update, 'validate_replenishment_result', lambda *a: events.append('validated'))
    monkeypatch.setattr(update, 'refresh_clearance_audit', lambda *a: events.append('audit'))
    monkeypatch.setattr(update, 'log_task', lambda *a: None)
    schemas = update.SchemaConfig(target_schema='etl_datasync_test', etl_source_schema='etl_datasync',
                                  dwd_source_schema='dwd_datasync', pricing_source_schema='temporary_dwd')
    update.execute_sql_step(conn, schemas, update.STEPS['replenishment_result'],
                            {'snapshot_date': date(2026, 10, 15), 'biz_date': date(2026, 10, 14)})
    assert events == ['validated', 'audit']
