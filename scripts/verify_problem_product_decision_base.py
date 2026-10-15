"""Validate the exported read-only decision workbook against its source extract."""

from __future__ import annotations

import json
import traceback
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "01a08f9c-f540-7580-8637-c21961e3c79c"
SOURCE = OUT / "problem_product_decision_base.json"
BOOK = OUT / "2026-09-14-七国问题产品逐条决策模拟底表.xlsx"
LOG = OUT / "problem_product_decision_base_validation.json"


def close(actual, expected, tolerance=0.011):
    return actual is not None and expected is not None and abs(float(actual) - float(expected)) <= tolerance


def main():
    source = json.loads(SOURCE.read_text(encoding="utf-8"))
    book = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    required = ["说明与汇总", "MSKU底表", "竞价对象", "搜索词候选", "名单差异"]
    assert book.sheetnames == required, book.sheetnames
    products = list(book["MSKU底表"].values)
    bids = list(book["竞价对象"].values)
    terms = list(book["搜索词候选"].values)
    diff = list(book["名单差异"].values)
    assert len(products) - 1 == len(source["rows"]) == 2118
    assert len(bids) - 1 == len(source["ad_objects"]) == 375
    assert len(terms) - 1 == len(source["terms"]) == 332
    assert len(diff) - 1 == len(source["cohort_diff"]) == 195
    issues = []
    candidates = 0
    for i, (actual, expected) in enumerate(zip(products[1:], source["rows"]), 2):
        if (actual[0], actual[1], actual[2]) != (expected["country"], expected["store"], expected["msku"]):
            issues.append(f"key mismatch at row {i}")
        proposed = actual[12]
        if expected["proposed_price"] is None:
            if proposed not in (None, ""):
                issues.append(f"unexpected price at row {i}")
        elif not close(proposed, expected["proposed_price"]):
            issues.append(f"price mismatch at row {i}: {proposed} vs {expected['proposed_price']}")
        if proposed is not None and actual[9] is not None and float(proposed) < float(actual[9]) and actual[11] is not None and float(proposed) + 1e-6 < float(actual[11]):
            issues.append(f"floor violated at row {i}")
        if proposed is not None and actual[9] is not None and float(proposed) < float(actual[9]):
            candidates += 1
        if expected["monthly_balance_upper_cny"] is None:
            if actual[18] not in (None, ""):
                issues.append(f"unexpected month balance at row {i}")
        elif not close(actual[18], expected["monthly_balance_upper_cny"]):
            issues.append(f"month balance mismatch at row {i}")
    assert not issues, issues[:20]
    assert candidates == source["summary"]["price_reduction_candidates"] == 1033, candidates
    result = {"status": "passed", "rows": len(products) - 1, "price_candidates": candidates,
              "ad_objects": len(bids) - 1, "search_terms": len(terms) - 1,
              "cohort_diff": len(diff) - 1, "issues": 0}
    LOG.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception:
        LOG.write_text(json.dumps({"status": "failed", "traceback": traceback.format_exc()}, ensure_ascii=False), encoding="utf-8")
        raise
