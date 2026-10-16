"""Independent checks for the pending-review JSON and exported workbook."""

from __future__ import annotations

import json
from collections import Counter
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("outputs/问题产品广告待审核建议-20260914")
DATA = ROOT / "广告建议只读数据.json"
WORKBOOK = ROOT / "2026-09-13两组问题产品广告待审核建议.xlsx"


def main():
    data = json.loads(DATA.read_text(encoding="utf-8"))
    products, bids, budgets, terms = (data[name] for name in ("products", "bids", "budgets", "terms"))
    assert len(products) == 1866
    assert Counter(p["产品组"] for p in products) == {"无销量售价不高于P15": 1037, "日销0至1毛利低于5%": 829}
    product_keys = {(p["国家"], p["店铺"], p["MSKU"]) for p in products}
    assert len(product_keys) == 1866
    assert sum(Counter(p["广告状态"] for p in products).values()) == 1866
    assert len({(b["profile_id"], b["campaign_id"]) for b in budgets}) == len(budgets)
    for row in bids:
        new = row.get("new_bid")
        if new is None:
            continue
        new, old = Decimal(str(new)), Decimal(str(row["current_bid"]))
        if new < old:
            assert new >= Decimal("0.10")
        if new > old:
            assert new <= old * Decimal("1.20")
            assert new <= Decimal(str(row["max_bid"]))
        assert row["condition"].startswith(("降价", "提价"))
    for row in budgets:
        if row["shared_outside_cohort"]:
            assert row["new_budget"] is None
        if row["status"] in ("待复核", "暂停活动"):
            assert row["new_budget"] is None
        if row["new_budget"] is not None:
            assert Decimal(str(row["current_daily_budget"])) > 0
    for row in terms:
        assert not (row["加词建议"] == "候选加词" and row["否词建议"] == "候选否词")
        if row["否词建议"] == "候选否词":
            assert not row["受保护"]
    book = load_workbook(WORKBOOK, read_only=True, data_only=False)
    expected = {"产品状态": len(products), "竞价明细": len(bids), "预算明细": len(budgets), "加词否词明细": len(terms)}
    assert set(book.sheetnames) == {*expected, "条件汇总", "数据口径"}
    for name, count in expected.items():
        observed_rows = book[name].max_row or sum(1 for _ in book[name].iter_rows())
        assert observed_rows == count + 4, (name, observed_rows, count)
    sample = book["产品状态"]["C5"].value
    assert sample == products[0]["MSKU"]
    long_id_index = next(i for i, row in enumerate(bids, 5) if len(str(row["profile_id"])) > 15)
    assert book["竞价明细"].cell(long_id_index, 5).value == str(bids[long_id_index - 5]["profile_id"])
    print(json.dumps({"products": len(products), "product_groups": dict(Counter(p["产品组"] for p in products)), "statuses": dict(Counter(p["广告状态"] for p in products)), "bid_rows": len(bids), "bid_new_values": sum(b["new_bid"] is not None for b in bids), "budget_rows": len(budgets), "budget_new_values": sum(b["new_budget"] is not None for b in budgets), "term_rows": len(terms), "workbook_sheets": book.sheetnames}, ensure_ascii=False))


if __name__ == "__main__":
    main()
