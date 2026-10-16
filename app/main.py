import csv
import io
from contextlib import asynccontextmanager
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Literal, Optional
from urllib.parse import quote

from fastapi import Depends, FastAPI, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from etl.dashboard_daily_update import COLUMN_COMMENTS

from .services.dashboard_db import dashboard_service
from .services.ad_budget_data import ad_budget_service
from .services.ad_performance_data import ad_performance_service
from .services.sp_recommendation_data import sp_recommendation_service
from .services.sp_product_diagnosis_data import sp_product_diagnosis_service
from .services.sp_governance_data import sp_governance_service
from .services.country_label_hub_data import country_label_hub_service
from .services.label_hub_data import LabelHubRefreshBusyError, label_hub_service
from .services.label_hub_diagnostics import label_hub_diagnostics_service
from .services.label_hub_role_diagnostics import label_hub_role_diagnostic_service
from .services.label_hub_detail_data import label_hub_detail_service
from .services.label_hub_export import label_hub_export_service
from .services.label_hub_change_data import label_hub_change_service
from .services.price_review_data import price_review_service
from .services.replenishment_data import replenishment_service
from .services.replenishment_tracking_data import replenishment_tracking_service
from .services.replenishment_tracking_summary_data import replenishment_tracking_summary_service
from .services.return_goods_data import return_goods_service
from .services.return_goods_strategy_data import return_goods_strategy_service


BASE_DIR = Path(__file__).resolve().parent
TWO_DECIMAL_EXPORT_COLUMNS = {
    "limit_price",
    "limit_price_10",
    "limit_price_35",
    "margin_price_35",
    "margin_price_10",
}
CSV_FORMULA_PREFIXES = ("=", "+", "-", "@")
REPLENISHMENT_EXPORT_HEADER_COMMENTS = {
    "fba_local_quantity": "FBA本地 = FBA总库存+本地库存",
    "total": "FBA总库存 = FBA可售+FBA预留+待调仓+标发在途+入库中",
    "available_total": "FBA可用库存 = FBA可售+待调仓+FBA预留+入库中（此字段可在业务配置中自定义）",
    "afn_fulfillable_quantity": "FBA可售库存 = afn fulfillable",
    "stock_up_num": "FBA在途 = 实际在途发货单发货数量-签收数量",
    "afn_unsellable_quantity": "FBA不可售 = Unfulfillable",
    "sc_quantity_local_valid": (
        "本地可用 = 配对SKU的可用量+可用锁定量+期望可用量，"
        "仅组合产品包含期望可用量；点击设置-业务配置-补货建议中自定义"
    ),
    "sc_quantity_purchase_shipping": "采购在途 = 目的仓为本地仓的调拨单待收货量",
    "sc_quantity_purchase_plan": (
        "采购计划 = 配对SKU相关采购计划单待采购量统计数据："
        "采购计划（待审批）+采购计划（待采购）"
    ),
    "sc_quantity_local_qc": (
        "本地质检 = 配对SKU的待检待上架量（汇总SKU无绑定FNSKU数量与SKU+FNSKU数量）"
    ),
    "local_quantity": "本地库存 = 本地可用+采购在途+采购计划+本地质检",
    "export_annotations": "逐项列出本行实际触发的提醒；颜色、边框及触发条件详见最后的“颜色与标注说明”Sheet。标注不修改补货数量。",
    "final_sales_3d": (
        "黄色提醒包括：1. 当前3天销量集中；2. 最新完整销售日疑似单日爆单；"
        "3. 近30天内已确认爆单后恢复正常，但异常销量仍影响当前补货日销。"
        "黄色仅用于人工复核，不修改补货数量。"
    ),
}
REPLENISHMENT_SALES_CONCENTRATION_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)
REPLENISHMENT_DISABLED_STORE_FILL = PatternFill(
    fill_type="solid",
    fgColor="F4CCCC",
)
REPLENISHMENT_COMBINED_WARNING_FILL = PatternFill(
    fill_type="solid",
    fgColor="F4B183",
)
REPLENISHMENT_FOLLOW_FILL = PatternFill(
    fill_type="solid",
    fgColor="DDEBF7",
)
REPLENISHMENT_FOLLOW_BORDER_SIDE = Side(
    style="thick",
    color="4472C4",
)
REPLENISHMENT_EXPORT_HIDDEN_COLUMNS = {
    "onsale_sites",
    "unsale_sites",
    "marketplace_concat",
    "fba_local_quantity",
    "stock_up_num",
    "pre_daily_avg_sales",
    "pre_normal_replenish_need_qty",
    "pre_replenish_trigger_qty",
    "hist_90d_instock_days",
    "hist_90d_instock_sales",
    "hist_90d_instock_daily_sales",
    "history_recovery_need_qty",
    "history_recovery_flag",
    "support_inventory_qty",
    "predict_abcd_category",
    "r_90d_salable_days",
    "sales_180d",
    "sales_90d",
    "amount_180d",
    "amount_90d",
    "pprofit_180d",
    "pprofit_90d",
    "pprofit_ratio_180d",
    "pprofit_ratio_90d",
    "60d_stocko_qty",
    "90d_stocko_qty",
    "180d_stocko_qty",
    "amz_instock_sales_ratio",
    "instock_intrans_pur_sales_ratio",
    "purchase_lead_days_raw",
    "effective_purchase_lead_days",
    "purchase_lead_status",
    "arrival_inventory_support_days",
    "arrival_inventory_qty",
    "lead_time_demand_qty",
    "lead_time_stockout_flag",
    "lead_time_stockout_days",
    "fllow_flag",
    "followed_flag",
    "followed_by_count",
    "followed_by_links",
    "follow_origin_link",
    "replenish_block_reason",
    "asin_merge_flag",
    "asin_merge_target",
    "asin_merge_reason",
    "supplier_moq",
    "moq_status",
    "moq_shortfall_qty",
    "created_at",
    "updated_at",
}
CSV_HEADER_LABELS = {
    "snapshot_date": "快照日期",
    "period_start": "周期开始",
    "period_end": "周期结束",
    "item_key": "商品键",
    "stat_period": "统计周期",
    "dt_year": "年",
    "dt_week": "周",
    "dt_month": "月",
    "dt_date": "日期",
    "seller_name_new": "店铺",
    "seller_name": "原店铺",
    "seller_sku_adj": "MSKU",
    "seller_sku": "卖家SKU",
    "country_category": "国家分组",
    "country": "国家",
    "local_sku": "SKU",
    "sales_qty": "销量",
    "sales_amount": "销售额",
    "sales_amount_ex_tax": "不含税销售额",
    "order_gross_profit": "毛利润",
    "raw_order_gross_profit": "原始毛利润",
    "order_gross_margin": "毛利率",
    "margin_band": "毛利分层",
    "settlement_gross_profit": "结算毛利润",
    "filter_flag": "纳入筛选",
    "current_price_cny": "当前售价(CNY)",
    "price_currency": "币种",
    "current_price": "当前售价",
    "limit_price": "限价",
    "limit_price_10": "10毛利定价",
    "limit_price_adj": "调整后限价",
    "limit_price_without_ad": "不含广告限价",
    "over_limit_flag": "是否超限价",
    "shipping_method": "发货方式",
    "target_margin": "目标毛利率",
    "in_stock_days": "有货天数",
    "stat_days": "统计天数",
    "abnormal_days": "异常天数",
    "all_abnormal_flag": "全周期异常",
    "abnormal_flag_count": "异常标记次数",
    "daily_sales": "日销",
    "daily_sales_in_stock_days": "有货日销",
    "daily_sales_in_stock_band": "有货日销分层",
    "daily_sales_band": "日销分层",
    "ad_spend": "广告花费",
    "ad_orders": "广告订单量",
    "ad_sales": "广告销售额",
    "ad_clicks": "广告点击量",
    "ad_impressions": "广告曝光量",
    "acos": "ACOS",
    "tacos": "TACOS",
    "ctr": "CTR",
    "sessions_total": "Sessions",
    "ranking": "排名",
    "return_count": "退货量",
    "return_amount": "退货额",
    "net_amount": "净销售额",
    "afn_fulfillable_quantity": "AFN可售库存",
    "fba_total_inventory": "FBA总库存",
    "fba_total_inventory_cost": "FBA总库存成本",
    "fba_available_inventory": "FBA可用库存",
    "fba_available_inventory_cost": "FBA可用库存成本",
    "fba_sellable_inventory": "FBA可售库存",
    "pending_transfer": "预留调拨",
    "transferring_qty": "调拨在途",
    "pending_shipment": "待发货",
    "unsellable_inventory": "不可售库存",
    "planned_inbound": "计划入库",
    "actual_in_transit": "实际在途",
    "under_investigation": "调查中库存",
    "total_available_inventory": "总可用库存",
    "local_sellable_inventory": "本地可售库存",
    "local_stock_sellable_days": "本地可售天数",
    "price": "当前售价",
    "org_currency_icon": "币种",
    "price_cny": "售价(CNY)",
    "tax_inclusive_price": "含税限价",
    "tax_inclusive_price_noad": "不含广告限价",
    "tax_inclusive_price_adj": "调整后限价",
    "margin_price_35": "35毛利定价",
    "margin_price_10": "10毛利定价",
    "created_at": "创建时间",
    "updated_at": "更新时间",
}

@asynccontextmanager
async def lifespan(_: FastAPI):
    ad_budget_service.warm_cache()
    yield


app = FastAPI(title="产品分层看板", version="1.0.0", lifespan=lifespan)

from .label_analysis_routes import router as label_analysis_router
app.include_router(label_analysis_router)
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))


class AdjustmentDayNotePayload(BaseModel):
    note: str = ""


class LabelHubDetailRequest(BaseModel):
    detail_view: Literal["business_unit", "country"] = "business_unit"
    data_date: str = ""
    metric_period: str = "30d"
    country_category: str = "all"
    store: str = "all"
    keyword: str = ""
    parent_label_id: int = 0
    compare_parent_id: int = 0
    conditions: str = ""
    label_period: str = "all"
    analysis_parent_ids: list[int] = Field(default_factory=list)
    analysis_periods: list[str] = Field(default_factory=list)
    identifiers: list[str] = Field(default_factory=list)
    country_categories: list[str] = Field(default_factory=list)
    stores: list[str] = Field(default_factory=list)
    countries: list[str] = Field(default_factory=list)
    sales_roles: list[str] = Field(default_factory=list)
    role_reason_ids: list[int] = Field(default_factory=list)
    current_stockout_only: bool = False
    stockout_before_role_period: str = ""
    stockout_before_role_ids: list[int] = Field(default_factory=list)
    stockout_operating_status_period: str = ""
    stockout_operating_status: str = ""
    stockout_insufficient_reason: str = ""
    stockout_operating_scope: Literal["", "business_unit", "country"] = ""
    stockout_operating_trend: str = ""
    stockout_history_dimension: str = ""
    stockout_history_code: str = ""
    stockout_history_period: str = ""
    daily_sales_bands: list[str] = Field(default_factory=list)
    margin_bands: list[str] = Field(default_factory=list)
    ranking_bands: list[str] = Field(default_factory=list)
    problems: list[str] = Field(default_factory=list)
    detail_conditions: str = ""
    problem_mode: Literal["any", "all"] = "any"
    page: int = Field(default=1, ge=1)
    page_size: Literal[20, 50, 100] = 20
    sort_field: str = "sales_amount"
    sort_dir: Literal["asc", "desc"] = "desc"


def _parse_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


def csv_cell_value(value, column: str | None = None):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if column in TWO_DECIMAL_EXPORT_COLUMNS:
        return f"{Decimal(str(value)):.2f}"
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, str):
        stripped = value.lstrip()
        if stripped.startswith(CSV_FORMULA_PREFIXES):
            return "\t" + value
    return value


def csv_header_value(column: str) -> str:
    return COLUMN_COMMENTS.get(column) or CSV_HEADER_LABELS.get(column, column)


def xlsx_cell_value(value, column: str | None = None):
    if isinstance(value, Decimal):
        return value
    return csv_cell_value(value, column)


def is_replenishment_sales_concentrated(row: dict) -> bool:
    try:
        level_sort = int(row.get("support_replenish_level_sort") or 0)
        sales_3d = Decimal(str(row.get("final_sales_3d") or 0))
        sales_7d = Decimal(str(row.get("final_sales_7d") or 0))
        replenish_qty = Decimal(str(row.get("replenish_qty") or 0))
    except (ArithmeticError, TypeError, ValueError):
        return False

    if level_sort not in (1, 2, 3):
        return False
    if replenish_qty <= 0:
        return False
    sales_spike_flag = str(row.get("sales_spike_flag") or 0).strip().lower()
    if sales_spike_flag in {"1", "1.0", "true", "yes", "是"}:
        return True
    if sales_7d < 10:
        return False
    return sales_3d * 10 >= sales_7d * 7


def is_disabled_replenishment_store_with_qty(row: dict, disabled_stores: set[str]) -> bool:
    store = str(row.get("seller_name_new") or "").strip().casefold()
    if not store or store not in disabled_stores:
        return False
    try:
        return Decimal(str(row.get("replenish_qty") or 0)) > 0
    except (ArithmeticError, TypeError, ValueError):
        return False


def is_follow_replenishment_with_qty(row: dict) -> bool:
    follow_value = row.get("fllow_flag")
    if follow_value is None or isinstance(follow_value, bool):
        return False
    if isinstance(follow_value, str):
        is_follow = follow_value.strip().casefold() in {"是", "yes", "true", "0", "0.0"}
    else:
        try:
            is_follow = Decimal(str(follow_value)) == 0
        except (ArithmeticError, TypeError, ValueError):
            return False
    if not is_follow:
        return False
    try:
        return Decimal(str(row.get("replenish_qty") or 0)) > 0
    except (ArithmeticError, TypeError, ValueError):
        return False


def add_replenishment_export_legend(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("颜色与标注说明")
    sheet.append(["补货导出 · 颜色与标注说明"])
    sheet.merge_cells("A1:E1")
    sheet["A1"].font = Font(name="微软雅黑", size=16, bold=True, color="1F4E78")
    sheet.row_dimensions[1].height = 32
    sheet.append(["清货名单强制拦截，补货数量、箱数及货值置零；其他颜色和标注仅供复核，不调整补货数量，无提醒不代表审核通过。"])
    sheet.merge_cells("A2:E2")
    sheet.row_dimensions[2].height = 30
    sheet.append(["样式示例", "标注名称", "出现位置", "含义与触发条件", "阅读与处理建议"])
    entries = [
        ("黄色底色", "销量异常，请复核", "补货计划整行", "紧急/建议/计划补货且数量＞0；近7天销量≥10且近3天占比≥70%，或命中系统单日爆单提醒。", "复核销量集中或爆单对日销的影响；黄色不代表已确认销量造假。", REPLENISHMENT_SALES_CONCENTRATION_FILL, None),
        ("红色底色", "禁补店铺仍有补货量", "补货计划整行", "店铺配置为不补货，但本次导出补货数量＞0。", "核对店铺补货资格及承接链接；导出提示不自行清零。", REPLENISHMENT_DISABLED_STORE_FILL, None),
        ("橙色底色", "销量异常＋禁补店铺", "补货计划整行", "同时触发黄色与红色条件，整行使用橙色。", "两项原因均需复核，不是另一个独立等级。", REPLENISHMENT_COMBINED_WARNING_FILL, None),
        ("浅蓝底色", "跟卖补货", "补货计划整行", "该行是跟卖链接且补货数量＞0，未触发上述黄/红/橙预警。", "识别实际跟卖补货对象；同时有其他预警时保留预警底色。", REPLENISHMENT_FOLLOW_FILL, None),
        ("深蓝粗边框", "跟卖补货", "补货计划整行外框", "该行是跟卖链接且补货数量＞0，始终叠加深蓝框。", "蓝框表示跟卖属性，可与黄/红/橙底色同时出现。", None, None),
        ("绿色文字", "低毛利，有满足条件的国家", "低毛利核查两列", "总角色为问题或瘦狗，且总日均销量＞1件/天、总毛利率＜5%；当前类别存在国家日均销量＞1件/天且毛利率＞5%。", "查看国家列的具体日销和毛利率；绿色不表示产品整体健康或审核通过。", None, "375623"),
        ("红色文字", "低毛利，无满足条件的国家", "低毛利核查两列", "满足同一总维度核查前提，但当前类别没有国家同时满足日均销量＞1件/天且毛利率＞5%。", "结合具体国家利润复核；没有国家记录按未经营跳过，不提示缺失国家。", None, "9C0006"),
        ("换行文字", "导出标注", "补货计划最后一列", "逐项列出本行实际触发的销量异常、禁补店铺、跟卖及低毛利说明，可同时存在。", "先读文字原因，再看底色或边框；此列始终保留，不受业务列选择影响。", None, None),
        ("空白", "未触发该项提醒", "导出标注/低毛利核查", "仅表示本行未触发对应提醒；明星、潜力及其他未满足低毛利前提的行，低毛利两列留空。", "无底色或标注不等于审核通过，也不等于数据完整。", None, None),
        ("深蓝表头", "栏目标题", "各Sheet的表头", "仅用于区分列标题，不具有业务预警含义。", "可筛选补货计划表头；首行与前两列已冻结。", PatternFill(fill_type="solid", fgColor="1F4E78"), "FFFFFF"),
        ("文字提示", "库容已超占用/无剩余", "库容预警状态及说明", "直接读取店铺sid对应的当月仓储类型额度：来源剩余＜0为已超占用，剩余＝0为无剩余。四种销售角色均参与，不受低毛利核查前提限制。", "只提示，不清零、不压减、不拦截；商品类型仅为历史参考时单独说明，不隐藏已读取的店铺额度。", None, None),
        ("独立明细", "清货名单，禁止补货", "清货拦截明细（仅有正数需求时出现）", "仅列当前快照、国家类别、店铺及MSKU关键词范围内有正数参考需求的清货对象，不受角色、分层及MOQ筛选影响。参考需求量为拦截前数量，不可执行；其他链接可能承接需求，不代表整组补货净减少量。", "清货商品在补货计划中的数量、箱数及货值直接置零，旧快照导出同样拦截。明细保留拦截前需求量供核对。无正数记录不生成此Sheet。", None, None),
        ("文字提示", "本行补货超库容", "本行库容核查/本行预计超出", "本行体积＝补货件数×单件m³；本行体积＞所选类型有效剩余库容时预警，超出量＝max(0，本行体积－剩余)。相等不判超额。", "实际类型优先取FBA库存明细.仓储类型；缺失时按月仓储费分类估算并注明来源。尺寸优先最新Listing，其次月费表，最后SKU包装尺寸。只预警，不修改数量。", None, None),
        ("文字提示", "本行未超，店铺合计超库容", "店铺同类型合计核查/店铺合计预计超出", "本行未超，但同账号、国家类别、仓储类型的完整补货需求合计体积＞剩余库容。合计覆盖所选日期所有有效补货对象，不受当前分页、角色周期或筛选影响。", "不同国家额度不相加；没有有效代表站点且各站点额度有差异时，按当前类别最小有效剩余值估算，完整站点值保留。", None, None),
        ("文字提示", "本次补货未超库容", "库容预警状态/体积核查说明", "本行和店铺同类型补货需求合计均≤有效剩余库容，且对应需求的类型、尺寸和额度数据完整。默认补货当月开始发货流程。", "这是本次需求全部新增发货的体积预警估算，不是发货许可；平台剩余已扣在库与开放货件，本次不再重复扣除。", None, None),
        ("文字提示", "本行未超，合计数据不足/无法核算", "库容预警状态/体积核查说明", "能计算本行但其他正数补货对象缺类型或体积时，不宣称整组放得下；已知体积已超额度则仍预警。确实缺类型、体积、额度或来源日期不符时，具体说明缺项。", "不再统一要求手动确认商品类型；缺失不写成零。数量为0时显示本行无需补货，新增体积为0。", None, None),
        ("文字提示", "来源有剩余库容", "库容预警状态", "合法确定额度、采集日期符合所选月份的读取规则，来源剩余＞0。没有共享组记录、商品尺寸或类型未确认，都不影响展示已有额度。", "仅描述来源额度有剩余，不表示本次补货放得下；是否可以比较体积，见补货体积核查说明。", None, None),
        ("文字提示", "已读取库容（站点有差异）", "店铺站点库容明细/匹配依据", "同店铺当前类别多个站点额度有差异，且没有可用的Standard共享组代表站点关联。逐站显示额度，不求和，不任意选择一个站点作为共享额度。", "单一剩余量留空；同值站点仅合并显示，不据此认定共享关系。其他店铺或加拿大、墨西哥额度不会混入。", None, None),
        ("文字提示", "商品类型未对应/店铺当月额度未提供", "店铺站点库容明细/补货体积核查说明", "商品类型未对应时仍列出已读取的各类型店铺额度；对应sid、月份或类型确实无记录时才说明额度未提供。", "标准件、大件、服装、鞋靴额度分开，不能混用或把缺失写成零。", None, None),
        ("文字提示", "预计库容不足", "库容预警状态及说明", "默认生成补货当月开始发货流程；整组新增需求体积完整且已核对平台占用时，将同期新增需求与剩余额度比较；需求＞剩余为不足。", "剩余额度已扣在库和开放货件，不能再次扣全部在途；确认发货月份不代表已核对整组需求和占用。", None, None),
        ("文字提示", "同期需求未超额", "库容预警状态及说明", "满足同一比较前提，整组同期新增需求≤剩余额度；恰好相等不判不足。", "仅指本次估算未超额，不表示平台许可或补货审核通过。", None, None),
        ("文字提示", "库容待核", "库容预警状态及说明", "当前月读取该月最新DWD版本，接受当天或前一天采集；历史月读取该月月末归档。采集日期和平台更新时间分别展示。采集异常、额度“－”或单位未知时说明具体来源问题。", "当前月只保存最新版本，历史月按月底归档，都不能宣称复原所选历史日的实时库容。默认生成补货当月发货；商品类型、尺寸、整组需求与货件去重问题单独列在体积核查说明。", None, None),
    ]
    for index, (sample, label, location, meaning, action, fill, color) in enumerate(entries, start=4):
        sheet.append([sample, label, location, meaning, action])
        sheet.row_dimensions[index].height = 64
        if fill:
            sheet.cell(index, 1).fill = fill
        if color:
            sheet.cell(index, 1).font = Font(name="微软雅黑", size=11, bold=True, color=color)
        if sample == "深蓝粗边框":
            sheet.cell(index, 1).border = Border(left=REPLENISHMENT_FOLLOW_BORDER_SIDE, right=REPLENISHMENT_FOLLOW_BORDER_SIDE,
                                               top=REPLENISHMENT_FOLLOW_BORDER_SIDE, bottom=REPLENISHMENT_FOLLOW_BORDER_SIDE)
    sheet.append(["国家核查范围：欧洲查德国、法国、意大利、西班牙、荷兰；英国查英国；北美查美国。同快照、同周期、同店铺＋MSKU，不跨国家类别。日销＝周期销量÷固定周期天数，毛利率＝毛利润合计÷销售额合计；日销恰好1或毛利率恰好5%不满足国家条件。"])
    sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=5)
    sheet.row_dimensions[sheet.max_row].height = 58
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    for cell in sheet[3]:
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    for letter, width in zip("ABCDE", (18, 26, 25, 70, 48)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A4"
    sheet.sheet_view.showGridLines = False


def build_replenishment_xlsx(payload: dict, disabled_stores: set[str] | None = None) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "补货计划"
    columns = [*payload["columns"], {"name": "export_annotations", "label": "导出标注"}]
    review_colors = {"qualified": "375623", "none": "9C0006", "insufficient": "666666"}
    normal_alignment = Alignment(vertical="center", wrap_text=False)
    annotation_alignment = Alignment(vertical="center", wrap_text=True)
    detail_alignment = Alignment(wrap_text=True, vertical="top")
    review_fonts = {status: Font(name="微软雅黑", size=11, bold=True, color=color)
                    for status, color in review_colors.items()}
    follow_borders = {
        column_index: Border(
            left=REPLENISHMENT_FOLLOW_BORDER_SIDE if column_index == 1 else Side(),
            right=REPLENISHMENT_FOLLOW_BORDER_SIDE if column_index == len(columns) else Side(),
            top=REPLENISHMENT_FOLLOW_BORDER_SIDE, bottom=REPLENISHMENT_FOLLOW_BORDER_SIDE,
        ) for column_index in range(1, len(columns) + 1)
    }
    normalized_disabled_stores = {
        str(store).strip().casefold()
        for store in (disabled_stores or set())
        if str(store).strip()
    }

    for column_index, column in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=column["label"])
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        width = 16
        if column["name"] in {"seller_sku_adj", "max_sku", "max_asin", "seller_name_new"}:
            width = 24
        elif column["name"] in {"low_margin_review_note", "qualifying_country_metrics", "export_annotations"}:
            width = 55 if column["name"] != "export_annotations" else 42
        elif column["name"] in {"capacity_note", "capacity_quota_summary", "capacity_product_note", "capacity_match_basis"}:
            width = 70
        elif column["name"] in {"capacity_group", "capacity_storage_type", "capacity_source_time", "capacity_measurement_source", "capacity_status_label"}:
            width = 32
        elif column["name"].startswith("capacity_"):
            width = 24
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
        if column["name"] in REPLENISHMENT_EXPORT_HIDDEN_COLUMNS:
            worksheet.column_dimensions[get_column_letter(column_index)].hidden = True
        comment_text = REPLENISHMENT_EXPORT_HEADER_COMMENTS.get(column["name"])
        if comment_text:
            cell.comment = Comment(comment_text, "看板系统")

    for row_index, row in enumerate(payload["rows"], start=2):
        highlight_row = is_replenishment_sales_concentrated(row)
        disabled_store_row = is_disabled_replenishment_store_with_qty(
            row,
            normalized_disabled_stores,
        )
        follow_replenishment_row = is_follow_replenishment_with_qty(row)
        annotations = []
        if row.get("replenish_block_reason") == "清货名单，禁止补货":
            annotations.append("清货名单，补货数量已置零")
        if highlight_row:
            annotations.append("销量异常，请复核")
        if disabled_store_row:
            annotations.append("禁补店铺仍有补货量")
        if follow_replenishment_row:
            annotations.append("跟卖补货（深蓝框）")
        if row.get("low_margin_review_status") in {"qualified", "none"}:
            country_state = "有国家" if row["low_margin_review_status"] == "qualified" else "无国家"
            annotations.append(f"总维度低毛利；{country_state}日均销量＞1件/天且毛利率＞5%")
        if row.get("capacity_status_label"):
            annotations.append(str(row["capacity_status_label"]) + "（见库容说明）")
        annotation_text = "\n".join(annotations)
        worksheet.row_dimensions[row_index].height = max(24, sum(max(1, (len(item) + 20) // 21) for item in annotations) * 16)
        if row.get("low_margin_review_status"):
            review_texts = [str(row.get(name) or "") for name in ("low_margin_review_note", "qualifying_country_metrics")]
            wrapped_lines = max((sum(max(1, (len(part) + 27) // 28) for part in value.splitlines()) for value in review_texts if value), default=1)
            worksheet.row_dimensions[row_index].height = max(worksheet.row_dimensions[row_index].height, wrapped_lines * 16)
        row_fill = None
        if highlight_row and disabled_store_row:
            row_fill = REPLENISHMENT_COMBINED_WARNING_FILL
        elif disabled_store_row:
            row_fill = REPLENISHMENT_DISABLED_STORE_FILL
        elif highlight_row:
            row_fill = REPLENISHMENT_SALES_CONCENTRATION_FILL
        elif follow_replenishment_row:
            row_fill = REPLENISHMENT_FOLLOW_FILL
        for column_index, column in enumerate(columns, start=1):
            value = annotation_text if column["name"] == "export_annotations" else xlsx_cell_value(row.get(column["name"]), column["name"])
            if column["name"] == "low_margin_review_note" and isinstance(value, str):
                value = value.replace("；", "；\n").replace("站点日均销量", "站点\n日均销量").replace("同时满足：", "同时满足：\n")
            elif column["name"] == "capacity_note" and isinstance(value, str):
                value = value.replace("；", "；\n")
            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )
            cell.alignment = annotation_alignment if column["name"] == "export_annotations" else normal_alignment
            if column["name"] in TWO_DECIMAL_EXPORT_COLUMNS and cell.value != "":
                cell.number_format = "0.00"
            if row_fill:
                cell.fill = row_fill
            if column["name"].startswith("capacity_"):
                if column["name"].endswith("_m3"):
                    cell.number_format = "0.00000000" if column["name"] == "capacity_unit_volume_m3" else "0.000000"
                else:
                    cell.alignment = detail_alignment
                    if value:
                        chars = max(1, int(worksheet.column_dimensions[get_column_letter(column_index)].width / 2))
                        lines = sum(max(1, (len(part) + chars - 1) // chars) for part in str(value).splitlines())
                        worksheet.row_dimensions[row_index].height = max(worksheet.row_dimensions[row_index].height, lines * 16 + 8)
            if column["name"] in {"low_margin_review_note", "qualifying_country_metrics"}:
                review_color = review_colors.get(row.get("low_margin_review_status"))
                if review_color:
                    cell.font = review_fonts[row["low_margin_review_status"]]
                cell.alignment = detail_alignment
                worksheet.column_dimensions[get_column_letter(column_index)].width = 55
                if value:
                    lines = sum(max(1, (len(part) + 27) // 28) for part in str(value).splitlines())
                    worksheet.row_dimensions[row_index].height = max(worksheet.row_dimensions[row_index].height, lines * 18 + 8)
            if follow_replenishment_row:
                cell.border = follow_borders[column_index]

    worksheet.row_dimensions[1].height = 44
    worksheet.freeze_panes = "C2" if len(columns) > 2 else "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    add_replenishment_clearance_sheet(workbook, payload.get("clearance_rows", []))
    add_replenishment_export_legend(workbook)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def add_replenishment_clearance_sheet(workbook: Workbook, rows: list[dict]) -> None:
    """A five-column appendix only for positive demand; keep the main sheet untouched."""
    positive = [row for row in rows if Decimal(str(row.get("reference_qty") or 0)) > 0]
    if not positive:
        return
    sheet = workbook.create_sheet("清货拦截明细")
    fields = ("country_category", "seller_name_new", "seller_sku_adj", "reference_qty", "block_reason")
    sheet.append(["国家类别", "店铺", "MSKU", "参考需求量（件）", "拦截原因"])
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill, cell.font = header_fill, header_font
    sheet["D1"].comment = Comment("参考需求量为清货拦截前数量，仅供核对，不可执行。补货计划中的清货商品数量、箱数及货值已置零。参考量沿用原ASIN承接、跟卖及MOQ口径，不代表整组补货净减少量。", "补货规则")
    for row in positive:
        sheet.append([xlsx_cell_value(row.get(field), field) for field in fields])
    for column, width in zip("ABCDE", (14, 20, 24, 22, 55)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions


def build_filters(
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    over_limit: str = Query(default="all"),
    daily_sales_band: str = Query(default="all"),
    margin_band: str = Query(default="all"),
    keyword: str = Query(default=""),
) -> dict:
    return {
        "start_date": start_date,
        "end_date": end_date,
        "site": site,
        "store": store,
        "over_limit": over_limit,
        "daily_sales_band": daily_sales_band,
        "margin_band": margin_band,
        "keyword": keyword,
    }


@app.get("/", response_class=HTMLResponse)
def dashboard_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "dashboard.html",
        {"page": "dashboard", "title": "产品分层看板"},
    )


@app.get("/layers", response_class=HTMLResponse)
def layers_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "layers.html",
        {"page": "layers", "title": "产品分层分布"},
    )


@app.get("/matrix", response_class=HTMLResponse)
def matrix_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "matrix.html",
        {"page": "matrix", "title": "日销毛利率矩阵"},
    )


@app.get("/detail", response_class=HTMLResponse)
def detail_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "detail.html",
        {"page": "detail", "title": "产品数据明细"},
    )


@app.get("/alerts", response_class=HTMLResponse)
def alerts_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "alerts.html",
        {"page": "alerts", "title": "异常预警工作台"},
    )


@app.get("/ad-budget", response_class=HTMLResponse)
def ad_budget_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "ad_budget.html",
        {"page": "ad_budget", "title": "广告预算执行工作台"},
    )


@app.get("/ad-performance", response_class=HTMLResponse)
def ad_performance_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "ad_performance.html",
        {"page": "ad_performance", "title": "SP 广告效果分析"},
    )


def _ad_performance_params(request: Request) -> dict:
    return dict(request.query_params)


@app.get("/api/ad-performance/options")
def api_ad_performance_options() -> dict:
    return ad_performance_service.options()


@app.get("/api/ad-performance/summary")
def api_ad_performance_summary(request: Request) -> dict:
    try:
        return ad_performance_service.summary(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.get("/api/ad-performance/trend")
def api_ad_performance_trend(request: Request) -> dict:
    try:
        return ad_performance_service.trend(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.get("/api/ad-performance/rows")
def api_ad_performance_rows(request: Request) -> dict:
    try:
        return ad_performance_service.rows(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.get("/api/ad-performance/detail")
def api_ad_performance_detail(request: Request) -> dict:
    try:
        return ad_performance_service.detail(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/export")
def api_ad_performance_export(request: Request) -> StreamingResponse:
    try:
        content = ad_performance_service.export_csv(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    return StreamingResponse(
        iter([content]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''sp-ad-performance.csv"},
    )


@app.get("/api/ad-performance/recommendations/options")
def api_sp_recommendation_options() -> dict:
    return sp_recommendation_service.options()


@app.get("/api/ad-performance/governance/options")
def api_sp_governance_options() -> dict:
    try:
        return sp_governance_service.options()
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/governance/summary")
def api_sp_governance_summary(request: Request) -> dict:
    try:
        return sp_governance_service.summary(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/governance/stores")
def api_sp_governance_stores(request: Request) -> dict:
    try:
        return sp_governance_service.stores(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/governance/entities")
def api_sp_governance_entities(request: Request) -> dict:
    try:
        return sp_governance_service.entities(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/recommendations/summary")
def api_sp_recommendation_summary(request: Request) -> dict:
    try:
        return sp_recommendation_service.summary(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/recommendations/rows")
def api_sp_recommendation_rows(request: Request) -> dict:
    try:
        return sp_recommendation_service.rows(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/recommendations/detail")
def api_sp_recommendation_detail(request: Request) -> dict:
    try:
        return sp_recommendation_service.detail(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/recommendations/export")
def api_sp_recommendation_export(request: Request) -> StreamingResponse:
    try:
        content = sp_recommendation_service.export_csv(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return StreamingResponse(
        iter([content]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''sp-ad-recommendations.csv"},
    )


@app.get("/api/ad-performance/product-diagnosis/options")
def api_sp_product_diagnosis_options() -> dict:
    return sp_product_diagnosis_service.options()


@app.get("/api/ad-performance/product-diagnosis/summary")
def api_sp_product_diagnosis_summary(request: Request) -> dict:
    try:
        return sp_product_diagnosis_service.summary(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/product-diagnosis/trend")
def api_sp_product_diagnosis_trend(request: Request) -> dict:
    try:
        return sp_product_diagnosis_service.trend(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/product-diagnosis/rows")
def api_sp_product_diagnosis_rows(request: Request) -> dict:
    try:
        return sp_product_diagnosis_service.rows(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/product-diagnosis/detail")
def api_sp_product_diagnosis_detail(request: Request) -> dict:
    try:
        return sp_product_diagnosis_service.detail(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/ad-performance/product-diagnosis/export")
def api_sp_product_diagnosis_export(request: Request) -> StreamingResponse:
    try:
        content = sp_product_diagnosis_service.export_csv(_ad_performance_params(request))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return StreamingResponse(iter([content]), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''sp-product-diagnosis.csv"})


@app.get("/api/ad-budget/meta")
def api_ad_budget_meta() -> dict:
    return ad_budget_service.get_meta()


@app.get("/api/ad-budget")
def api_ad_budget(
    country_category: str = "all",
    country: str = "all",
    seller_name_new: str = "all",
    product_type: str = "all",
    inventory_status: str = "all",
    anomaly: str = "all",
    keyword: str = "",
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=20, le=200),
    sort_field: str = "anomaly_priority",
    sort_dir: Literal["asc", "desc"] = "desc",
    column_filters: str = "",
) -> dict:
    return ad_budget_service.get_payload(
        country_category=country_category,
        country=country,
        seller_name_new=seller_name_new,
        product_type=product_type,
        inventory_status=inventory_status,
        anomaly=anomaly,
        keyword=keyword,
        page=page,
        page_size=page_size,
        sort_field=sort_field,
        sort_dir=sort_dir,
        column_filters=column_filters,
    )


@app.get("/api/ad-budget/detail")
def api_ad_budget_detail(
    country_category: str,
    country: str,
    seller_name_new: str,
    seller_sku_adj: str,
) -> dict:
    try:
        return ad_budget_service.get_detail(
            country_category=country_category,
            country=country,
            seller_name_new=seller_name_new,
            seller_sku_adj=seller_sku_adj,
        )
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/opportunities", response_class=HTMLResponse)
def opportunities_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "opportunities.html",
        {"page": "opportunities", "title": "机会 SKU 池"},
    )


@app.get("/inventory-weekly", response_class=HTMLResponse)
def inventory_weekly_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "inventory_weekly.html",
        {"page": "inventory_weekly", "title": "库存周报"},
    )


@app.get("/replenishment", response_class=HTMLResponse)
def replenishment_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "replenishment.html",
        {"page": "replenishment", "title": "补货计划"},
    )

@app.get("/replenishment-tracking", response_class=HTMLResponse)
def replenishment_tracking_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "replenishment_tracking.html",
        {"page": "replenishment_tracking", "title": "补货追踪"},
    )


@app.get("/return-goods", response_class=HTMLResponse)
def return_goods_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "return_goods.html",
        {"page": "return_goods", "title": "返场品"},
    )


@app.get("/return-goods-strategy", response_class=HTMLResponse)
def return_goods_strategy_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(request, "return_goods_strategy.html",
                                      {"page": "return_goods_strategy", "title": "返场品运营策略"})


@app.get("/api/return-goods-strategy")
def api_return_goods_strategy(
    snapshot_date: str = "", country_category: str = "all", seller_name_new: str = "all",
    keyword: str = "", exit_status: str = "tracking", state: str = "all", action: str = "all",
    recommendation: str = "all", stage: str = "all", age_band: str = "all",
    focus: str = "all", all_rows: bool = False,
    page: int = Query(1, ge=1), page_size: int = Query(50, ge=10, le=200),
):
    try:
        if snapshot_date:
            from datetime import date as strategy_date
            strategy_date.fromisoformat(snapshot_date)
        return return_goods_strategy_service.get_payload(
            snapshot_date=snapshot_date, country_category=country_category, seller_name_new=seller_name_new,
            keyword=keyword, exit_status=exit_status, state=state, action=action,
            recommendation=recommendation, stage=stage, age_band=age_band, focus=focus, all_rows=all_rows,
            page=page, page_size=page_size)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="日期或筛选参数无效") from exc


@app.get("/api/return-goods-strategy/detail")
def api_return_goods_strategy_detail(event_id: str, batch_id: int = Query(..., ge=1)):
    detail = return_goods_strategy_service.get_detail(event_id=event_id, batch_id=batch_id)
    if detail is None:
        raise HTTPException(status_code=404, detail="本批次没有该返场事件")
    return detail


@app.get("/label-hub", response_class=HTMLResponse)
def label_hub_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "label_hub.html",
        {"page": "label_hub", "title": "标签看板"},
    )


@app.get("/country-label-hub", response_class=HTMLResponse)
def country_label_hub_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "country_label_hub.html",
        {"page": "country_label_hub", "title": "国家标签看板"},
    )


@app.get("/api/label-hub/meta")
def api_label_hub_meta() -> dict:
    return label_hub_service.get_meta()


@app.post("/api/label-hub/refresh")
def api_label_hub_refresh() -> dict:
    try:
        return label_hub_service.force_source_refresh()
    except LabelHubRefreshBusyError as exc:
        raise HTTPException(status_code=409, detail="标签数据正在刷新，请稍后") from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="远端标签刷新失败，请稍后重试") from exc


@app.get("/api/label-hub/stockout-before-roles")
def api_label_hub_stockout_before_roles(
    data_date: str = "",
    role_period: str = "30d",
    country_category: str = "all",
    store: str = "all",
    keyword: str = "",
) -> dict:
    try:
        return label_hub_service.get_stockout_before_role_summary(
            data_date=data_date,
            role_period=role_period,
            country_category=country_category,
            store=store,
            keyword=keyword,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="断货前销售角色汇总暂不可用，请稍后重试") from exc


@app.get("/api/label-hub/stockout-operating-status")
def api_label_hub_stockout_operating_status(
    data_date: str = "",
    scope: Literal["business_unit", "country"] = "business_unit",
    country_category: str = "all",
    store: str = "all",
    keyword: str = "",
) -> dict:
    try:
        return label_hub_service.get_stockout_operating_status_summary(
            data_date=data_date,
            scope=scope,
            country_category=country_category,
            store=store,
            keyword=keyword,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="断货经营状态汇总暂不可用，请稍后重试") from exc


@app.get("/api/label-hub/stockout-before-role-evidence")
def api_label_hub_stockout_before_role_evidence(
    data_date: str,
    country_category: str,
    store: str,
    msku: str,
    role_period: str = "30d",
    country: str = "",
) -> dict:
    try:
        return label_hub_service.get_stockout_before_role_evidence(
            data_date=data_date,
            country_category=country_category,
            store=store,
            msku=msku,
            role_period=role_period,
            country=country,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="断货前销售角色依据暂不可用，请稍后重试") from exc


@app.get("/api/label-hub")
def api_label_hub(
    data_date: str = "", country_category: str = "all", store: str = "all", keyword: str = "",
    parent_label_id: int = 0, compare_parent_id: int = 0, conditions: str = "", label_period: str = "all",
    metric_period: str = "30d", analysis_parent_ids: str = "", analysis_periods: str = "",
    sales_roles: str = "", sales_trends: str = "", daily_sales_bands: str = "", margin_bands: str = "", problem: str = "all",
    page: int = 1, page_size: int = 20, sort_field: str = "sales_amount", sort_dir: str = "desc",
) -> dict:
    try:
        return label_hub_service.get_payload(
            data_date=data_date, country_category=country_category, store=store, keyword=keyword,
            parent_label_id=parent_label_id, compare_parent_id=compare_parent_id,
            conditions=conditions, label_period=label_period, metric_period=metric_period,
            analysis_parent_ids=analysis_parent_ids, analysis_periods=analysis_periods,
            sales_roles=sales_roles, sales_trends=sales_trends,
            daily_sales_bands=daily_sales_bands, margin_bands=margin_bands, problem=problem,
            page=page, page_size=page_size, sort_field=sort_field, sort_dir=sort_dir,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="标签数据暂不可用，请稍后重试") from exc


@app.get("/api/label-hub/msku")
def api_label_hub_msku(data_date: str, country_category: str, store: str, msku: str, metric_period: str = "30d") -> dict:
    try:
        return label_hub_service.get_msku_profile(data_date=data_date, country_category=country_category, store=store, msku=msku, metric_period=metric_period)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/label-hub/sales-role-diagnostics")
def api_label_hub_sales_role_diagnostics(
    data_date: str = "",
    metric_period: str = "30d",
    country_category: str = "all",
    store: str = "all",
    keyword: str = "",
    parent_label_id: int = 1,
    compare_parent_id: int = 2,
    conditions: str = "",
    label_period: str = "all",
    analysis_parent_ids: str = "",
    analysis_periods: str = "",
    sales_roles: str = "",
    sales_trends: str = "",
    daily_sales_bands: str = "",
    margin_bands: str = "",
    problem: str = "all",
    diagnostic_scope: str = "global",
    diagnostic_period: str = "30d",
) -> dict:
    try:
        return label_hub_diagnostics_service.get_payload(
            data_date=data_date,
            metric_period=metric_period,
            country_category=country_category,
            store=store,
            keyword=keyword,
            parent_label_id=parent_label_id,
            compare_parent_id=compare_parent_id,
            conditions=conditions,
            label_period=label_period,
            analysis_parent_ids=analysis_parent_ids,
            analysis_periods=analysis_periods,
            sales_roles=sales_roles,
            sales_trends=sales_trends,
            daily_sales_bands=daily_sales_bands,
            margin_bands=margin_bands,
            problem=problem,
            diagnostic_scope=diagnostic_scope,
            diagnostic_period=diagnostic_period,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="销售角色诊断暂不可用，请稍后重试") from exc


@app.get("/api/label-hub/msku-role-diagnostics")
def api_label_hub_msku_role_diagnostics(
    data_date: str,
    country_category: str,
    store: str,
    msku: str,
    diagnostic_period: str = "30d",
) -> dict:
    try:
        return label_hub_role_diagnostic_service.get_payload(
            data_date=data_date,
            country_category=country_category,
            store=store,
            msku=msku,
            diagnostic_period=diagnostic_period,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="商品角色诊断暂不可用，请稍后重试") from exc


@app.post("/api/label-hub/details")
def api_label_hub_details(request: LabelHubDetailRequest) -> dict:
    try:
        return label_hub_detail_service.get_details(**request.model_dump())
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="标签明细数据暂不可用，请稍后重试") from exc


@app.post("/api/label-hub/details/export")
def api_label_hub_details_export(request: LabelHubDetailRequest) -> StreamingResponse:
    try:
        filename, chunks = label_hub_export_service.build_export(**request.model_dump())
        return StreamingResponse(
            chunks,
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
            },
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="标签明细导出暂不可用，请稍后重试") from exc


@app.get("/api/label-hub/msku-country-profile")
def api_label_hub_msku_country_profile(
    country_category: str,
    store: str,
    msku: str,
    metric_period: str = "30d",
) -> dict:
    try:
        return country_label_hub_service.get_label_hub_country_profile(
            country_category=country_category,
            store=store,
            msku=msku,
            metric_period=metric_period,
        )
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="国家画像暂不可用，请稍后重试") from exc


@app.get("/api/label-hub/changes")
def api_label_hub_changes(
    data_date: str = "", country_category: str = "all", store: str = "all", keyword: str = "",
    parent_label_id: int = 0, compare_parent_id: int = 0, conditions: str = "", label_period: str = "all",
    metric_period: str = "30d", analysis_parent_ids: str = "", analysis_periods: str = "",
    sales_trends: str = "", daily_sales_bands: str = "", margin_bands: str = "", problem: str = "all",
    transition_period: str = "", change_type: str = "all", page: int = 1, page_size: int = 20,
    sort_field: str = "change_type", sort_dir: str = "asc",
    layer_change_parent: int = 0, layer_change_bucket: str = "", layer_change_period: str = "all",
    layer_transition_from: str = "", layer_transition_to: str = "",
    layer_transition_parent: int = 0, layer_transition_previous: str = "", layer_transition_current: str = "",
) -> dict:
    try:
        return label_hub_change_service.get_changes(
            data_date=data_date, country_category=country_category, store=store, keyword=keyword,
            parent_label_id=parent_label_id, compare_parent_id=compare_parent_id,
            conditions=conditions, label_period=label_period, metric_period=metric_period,
            analysis_parent_ids=analysis_parent_ids, analysis_periods=analysis_periods,
            sales_trends=sales_trends, daily_sales_bands=daily_sales_bands,
            margin_bands=margin_bands, problem=problem, transition_period=transition_period,
            change_type=change_type, page=page, page_size=page_size,
            sort_field=sort_field, sort_dir=sort_dir,
            layer_change_parent=layer_change_parent, layer_change_bucket=layer_change_bucket,
            layer_change_period=layer_change_period,
            layer_transition_from=layer_transition_from, layer_transition_to=layer_transition_to,
            layer_transition_parent=layer_transition_parent,
            layer_transition_previous=layer_transition_previous,
            layer_transition_current=layer_transition_current,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="标签变化数据暂不可用，请稍后重试") from exc


@app.get("/api/label-hub/msku-change")
def api_label_hub_msku_change(msku: str, transition_period: str = "30d") -> dict:
    try:
        return label_hub_change_service.get_msku_change(msku=msku, transition_period=transition_period)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/country-label-hub/meta")
def api_country_label_hub_meta() -> dict:
    try:
        return country_label_hub_service.get_meta()
    except Exception as exc:
        raise HTTPException(status_code=503, detail="国家标签元数据暂不可用，请稍后重试") from exc


@app.get("/api/country-label-hub")
def api_country_label_hub(
    data_date: str = "", metric_period: str = "30d", country_category: str = "all",
    store: str = "all", keyword: str = "", conditions: str = "", label_periods: str = "",
    sales_trends: str = "", daily_sales_bands: str = "", margin_bands: str = "",
    problem: str = "all", page: int = 1, page_size: int = 20,
    sort_field: str = "problem_priority", sort_dir: str = "desc",
) -> dict:
    try:
        return country_label_hub_service.get_payload(
            data_date=data_date, metric_period=metric_period, country_category=country_category,
            store=store, keyword=keyword, conditions=conditions, label_periods=label_periods,
            sales_trends=sales_trends, daily_sales_bands=daily_sales_bands,
            margin_bands=margin_bands, problem=problem, page=page, page_size=page_size,
            sort_field=sort_field, sort_dir=sort_dir,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="国家标签数据暂不可用，请稍后重试") from exc


@app.get("/api/country-label-hub/msku")
def api_country_label_hub_msku(
    data_date: str, country: str, country_category: str, store: str, msku: str, metric_period: str = "30d",
) -> dict:
    try:
        return country_label_hub_service.get_msku_profile(
            data_date=data_date, country=country, country_category=country_category, store=store,
            msku=msku, metric_period=metric_period,
        )
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="国家标签画像暂不可用，请稍后重试") from exc


@app.get("/api/meta")
def api_meta() -> dict:
    return dashboard_service.get_meta()


@app.get("/api/sales-role/meta")
def api_sales_role_meta() -> dict:
    return dashboard_service.get_sales_role_meta()


@app.get("/api/sales-role")
def api_sales_role(
    period: str = "30d",
    country_category: str = "all",
    seller_name_new: str = "all",
    sales_role: str = "all",
    daily_sales_band: str = "all",
    margin_band: str = "all",
    keyword: str = "",
    page: int = 1,
    page_size: int = 20,
    sort_field: str = "sales_amount",
    sort_dir: str = "desc",
) -> dict:
    return dashboard_service.get_sales_role_payload(
        period=period,
        country_category=country_category,
        seller_name_new=seller_name_new,
        sales_role=sales_role,
        daily_sales_band=daily_sales_band,
        margin_band=margin_band,
        keyword=keyword,
        page=page,
        page_size=page_size,
        sort_field=sort_field,
        sort_dir=sort_dir,
    )


@app.get("/api/sales-role/export")
def api_sales_role_export(
    period: str = "30d",
    country_category: str = "all",
    seller_name_new: str = "all",
    sales_role: str = "all",
    daily_sales_band: str = "all",
    margin_band: str = "all",
    keyword: str = "",
    sort_field: str = "sales_amount",
    sort_dir: str = "desc",
) -> StreamingResponse:
    payload = dashboard_service.get_sales_role_export_payload(
        period=period,
        country_category=country_category,
        seller_name_new=seller_name_new,
        sales_role=sales_role,
        daily_sales_band=daily_sales_band,
        margin_band=margin_band,
        keyword=keyword,
        sort_field=sort_field,
        sort_dir=sort_dir,
    )
    output = io.StringIO(newline="")
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow([
        "销售角色", "国家类别", "店铺新", "MSKU", "SKU示例", "覆盖国家数", "覆盖国家",
        "销量", "日均销量", "销售额", "订单毛利额", "订单毛利率", "日销分层", "毛利率分层",
        "广告花费", "广告销售额", "ACOS", "TACOS",
    ])
    for row in payload["rows"]:
        writer.writerow([
            csv_cell_value(row.get("sales_role")),
            csv_cell_value(row.get("country_category")),
            csv_cell_value(row.get("seller_name_new")),
            csv_cell_value(row.get("seller_sku_adj")),
            csv_cell_value(row.get("local_sku_sample")),
            csv_cell_value(row.get("country_count")),
            csv_cell_value(row.get("countries")),
            csv_cell_value(row.get("sales_qty")),
            csv_cell_value(row.get("daily_sales")),
            csv_cell_value(row.get("sales_amount")),
            csv_cell_value(row.get("order_gross_profit")),
            csv_cell_value(row.get("order_gross_margin")),
            csv_cell_value(row.get("daily_sales_band")),
            csv_cell_value(row.get("margin_band")),
            csv_cell_value(row.get("ad_spend")),
            csv_cell_value(row.get("ad_sales")),
            csv_cell_value(row.get("acos")),
            csv_cell_value(row.get("tacos")),
        ])
    filename = f"sales_role_{period}_{date.today().isoformat()}.csv"
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers=headers)


@app.get("/api/sales-role/lifecycle")
def api_sales_role_lifecycle(
    period: str = "30d",
    country_category: str = "all",
    seller_name_new: str = "all",
    lifecycle_label: str = "all",
    sales_role: str = "all",
    keyword: str = "",
    page: int = 1,
    page_size: int = 20,
    sort_field: str = "sales_amount",
    sort_dir: str = "desc",
) -> dict:
    return dashboard_service.get_sales_role_lifecycle_payload(
        period=period,
        country_category=country_category,
        seller_name_new=seller_name_new,
        lifecycle_label=lifecycle_label,
        sales_role=sales_role,
        keyword=keyword,
        page=page,
        page_size=page_size,
        sort_field=sort_field,
        sort_dir=sort_dir,
    )


@app.get("/api/sales-role/lifecycle/export")
def api_sales_role_lifecycle_export(
    period: str = "30d",
    country_category: str = "all",
    seller_name_new: str = "all",
    lifecycle_label: str = "all",
    sales_role: str = "all",
    keyword: str = "",
    sort_field: str = "sales_amount",
    sort_dir: str = "desc",
) -> StreamingResponse:
    payload = dashboard_service.get_sales_role_lifecycle_export_payload(
        period=period,
        country_category=country_category,
        seller_name_new=seller_name_new,
        lifecycle_label=lifecycle_label,
        sales_role=sales_role,
        keyword=keyword,
        sort_field=sort_field,
        sort_dir=sort_dir,
    )
    output = io.StringIO(newline="")
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow([
        "生命周期", "销售角色", "国家类别", "店铺", "MSKU", "SKU示例", "覆盖国家数", "覆盖国家",
        "销量", "日均销量", "销售额", "订单毛利润", "订单毛利率", "标签周期", "销售角色周期",
        "广告花费", "广告销售额", "ACOS", "TACOS",
    ])
    for row in payload["rows"]:
        writer.writerow([
            csv_cell_value(row.get("lifecycle_label")),
            csv_cell_value(row.get("sales_role")),
            csv_cell_value(row.get("country_category")),
            csv_cell_value(row.get("seller_name_new")),
            csv_cell_value(row.get("seller_sku_adj")),
            csv_cell_value(row.get("local_sku_sample")),
            csv_cell_value(row.get("country_count")),
            csv_cell_value(row.get("countries")),
            csv_cell_value(row.get("sales_qty")),
            csv_cell_value(row.get("daily_sales")),
            csv_cell_value(row.get("sales_amount")),
            csv_cell_value(row.get("order_gross_profit")),
            csv_cell_value(row.get("order_gross_margin")),
            csv_cell_value(row.get("label_period")),
            csv_cell_value(row.get("sales_role_period")),
            csv_cell_value(row.get("ad_spend")),
            csv_cell_value(row.get("ad_sales")),
            csv_cell_value(row.get("acos")),
            csv_cell_value(row.get("tacos")),
        ])
    filename = f"sales_role_lifecycle_{period}_{date.today().isoformat()}.csv"
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers=headers)


@app.get("/api/dashboard")
def api_dashboard(
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    over_limit: str = Query(default="all"),
    daily_sales_band: str = Query(default="all"),
    margin_band: str = Query(default="all"),
    keyword: str = Query(default=""),
) -> dict:
    filters = build_filters(
        start_date=start_date,
        end_date=end_date,
        site=site,
        store=store,
        over_limit=over_limit,
        daily_sales_band=daily_sales_band,
        margin_band=margin_band,
        keyword=keyword,
    )
    return dashboard_service.get_dashboard_payload(filters)


@app.get("/api/dashboard/monthly-goals")
def api_dashboard_monthly_goals() -> dict:
    return dashboard_service.get_monthly_goals_payload()


@app.get("/api/alerts")
def api_alerts(
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    over_limit: str = Query(default="all"),
    daily_sales_band: str = Query(default="all"),
    margin_band: str = Query(default="all"),
    keyword: str = Query(default=""),
    alert_type: str = Query(default="all"),
    compare_days: int = Query(default=7, ge=7, le=90),
    comparison_code: str = Query(default=""),
    comparison_mode: str = Query(default="days"),
    previous_month: str = Query(default=""),
    recent_month: str = Query(default=""),
    sales_trend: str = Query(default="all"),
    rank_trend: str = Query(default="all"),
    margin_status: str = Query(default="all"),
    stock_status: str = Query(default="all"),
    transition_filter: str = Query(default=""),
    sort_field: str = Query(default=""),
    sort_dir: str = Query(default=""),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=10, le=100),
) -> dict:
    filters = build_filters(
        start_date=start_date,
        end_date=end_date,
        site=site,
        store=store,
        over_limit=over_limit,
        daily_sales_band=daily_sales_band,
        margin_band=margin_band,
        keyword=keyword,
    )
    return dashboard_service.get_alerts_payload(
        filters,
        alert_type=alert_type,
        compare_days=compare_days,
        comparison_code=comparison_code,
        comparison_mode=comparison_mode,
        previous_month=previous_month,
        recent_month=recent_month,
        sales_trend=sales_trend,
        rank_trend=rank_trend,
        margin_status=margin_status,
        stock_status=stock_status,
        transition_filter=transition_filter,
        sort_field=sort_field,
        sort_dir=sort_dir,
        page=page,
        page_size=page_size,
    )


@app.get("/api/alerts/export")
def api_alerts_export(
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    over_limit: str = Query(default="all"),
    daily_sales_band: str = Query(default="all"),
    margin_band: str = Query(default="all"),
    keyword: str = Query(default=""),
    alert_type: str = Query(default="all"),
    compare_days: int = Query(default=7, ge=7, le=90),
    comparison_code: str = Query(default=""),
    comparison_mode: str = Query(default="days"),
    previous_month: str = Query(default=""),
    recent_month: str = Query(default=""),
    sales_trend: str = Query(default="all"),
    rank_trend: str = Query(default="all"),
    margin_status: str = Query(default="all"),
    stock_status: str = Query(default="all"),
    transition_filter: str = Query(default=""),
) -> StreamingResponse:
    filters = build_filters(
        start_date=start_date,
        end_date=end_date,
        site=site,
        store=store,
        over_limit=over_limit,
        daily_sales_band=daily_sales_band,
        margin_band=margin_band,
        keyword=keyword,
    )
    payload = dashboard_service.get_alerts_export_payload(
        filters,
        alert_type=alert_type,
        compare_days=compare_days,
        comparison_code=comparison_code,
        comparison_mode=comparison_mode,
        previous_month=previous_month,
        recent_month=recent_month,
        sales_trend=sales_trend,
        rank_trend=rank_trend,
        margin_status=margin_status,
        stock_status=stock_status,
        transition_filter=transition_filter,
    )

    output = io.StringIO(newline="")
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow([
        "统计周期", "对比周期", "类型", "MSKU", "店铺", "国家",
        "销量趋势", "近期总销量", "近期日均销量", "对比总销量", "对比日均销量",
        "近期销售额", "近期日均销售额", "对比销售额", "对比日均销售额",
        "排名趋势", "毛利", "库存",
    ])
    for item in payload["items"]:
        writer.writerow(
            [
                payload["window"],
                payload["comparison_window"],
                csv_cell_value(item.get("label")),
                csv_cell_value(item.get("title")),
                csv_cell_value(item.get("store")),
                csv_cell_value(item.get("country")),
                csv_cell_value(item.get("sales_text")),
                csv_cell_value(item.get("recent_qty")),
                csv_cell_value(item.get("recent_daily_sales")),
                csv_cell_value(item.get("previous_qty")),
                csv_cell_value(item.get("previous_daily_sales")),
                csv_cell_value(item.get("recent_sales_amount")),
                csv_cell_value(item.get("recent_daily_sales_amount")),
                csv_cell_value(item.get("previous_sales_amount")),
                csv_cell_value(item.get("previous_daily_sales_amount")),
                csv_cell_value(item.get("rank_text")),
                csv_cell_value(item.get("margin_text")),
                csv_cell_value(item.get("stock_text")),
            ]
        )

    filename = f"alerts_{payload.get('comparison_code') or str(payload['compare_days']) + 'd'}_{date.today().isoformat()}.csv"
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
    }
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers=headers)


@app.get("/api/opportunities")
def api_opportunities(
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    country: str = Query(default="all"),
    over_limit: str = Query(default="no"),
    keyword: str = Query(default=""),
    opportunity_type: str = Query(default="all"),
    compare_days: int = Query(default=14, ge=7, le=90),
    comparison_code: str = Query(default=""),
    comparison_mode: str = Query(default="days"),
    previous_month: str = Query(default=""),
    recent_month: str = Query(default=""),
    stock_status: str = Query(default="all"),
    transition_filter: str = Query(default=""),
    sort_field: str = Query(default=""),
    sort_dir: str = Query(default=""),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=10, le=100),
) -> dict:
    filters = build_filters(
        start_date=None,
        end_date=None,
        site=country if country != "all" else site,
        store=store,
        over_limit=over_limit,
        daily_sales_band="all",
        margin_band="all",
        keyword=keyword,
    )
    return dashboard_service.get_opportunities_payload(
        filters,
        opportunity_type=opportunity_type,
        compare_days=compare_days,
        comparison_code=comparison_code,
        comparison_mode=comparison_mode,
        previous_month=previous_month,
        recent_month=recent_month,
        stock_status=stock_status,
        transition_filter=transition_filter,
        sort_field=sort_field,
        sort_dir=sort_dir,
        page=page,
        page_size=page_size,
    )


@app.get("/api/opportunities/export")
def api_opportunities_export(
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    country: str = Query(default="all"),
    over_limit: str = Query(default="no"),
    keyword: str = Query(default=""),
    opportunity_type: str = Query(default="all"),
    compare_days: int = Query(default=14, ge=7, le=90),
    comparison_code: str = Query(default=""),
    comparison_mode: str = Query(default="days"),
    previous_month: str = Query(default=""),
    recent_month: str = Query(default=""),
    stock_status: str = Query(default="all"),
    transition_filter: str = Query(default=""),
) -> StreamingResponse:
    filters = build_filters(
        start_date=None,
        end_date=None,
        site=country if country != "all" else site,
        store=store,
        over_limit=over_limit,
        daily_sales_band="all",
        margin_band="all",
        keyword=keyword,
    )
    payload = dashboard_service.get_opportunities_export_payload(
        filters,
        opportunity_type=opportunity_type,
        compare_days=compare_days,
        comparison_code=comparison_code,
        comparison_mode=comparison_mode,
        previous_month=previous_month,
        recent_month=recent_month,
        stock_status=stock_status,
        transition_filter=transition_filter,
    )

    output = io.StringIO(newline="")
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow([
        "统计周期", "对比周期", "机会类型", "机会分", "MSKU", "店铺", "国家",
        "日销", "销量变化", "销售额", "毛利率", "毛利润", "排名变化",
        "Sessions", "转化率", "FBA可售", "可售天数", "ACOS", "TACOS",
        "广告花费", "售价", "35毛利定价", "10毛利定价", "是否超限价", "建议动作",
    ])
    for item in payload["items"]:
        writer.writerow([
            payload["window"],
            payload["comparison_window"],
            csv_cell_value(item.get("label")),
            csv_cell_value(item.get("score")),
            csv_cell_value(item.get("msku")),
            csv_cell_value(item.get("store")),
            csv_cell_value(item.get("country")),
            csv_cell_value(item.get("daily_sales")),
            csv_cell_value(item.get("sales_text")),
            csv_cell_value(item.get("scoped_revenue")),
            csv_cell_value(item.get("margin")),
            csv_cell_value(item.get("profit")),
            csv_cell_value(item.get("rank_text")),
            csv_cell_value(item.get("recent_sessions")),
            csv_cell_value(item.get("conversion")),
            csv_cell_value(item.get("fba_sellable_inventory")),
            csv_cell_value(item.get("sellable_days")),
            csv_cell_value(item.get("acos")),
            csv_cell_value(item.get("tacos")),
            csv_cell_value(item.get("ad_spend")),
            csv_cell_value(item.get("current_price")),
            csv_cell_value(item.get("limit_price_35")),
            csv_cell_value(item.get("limit_price_10")),
            "是" if item.get("over_limit") else "否",
            csv_cell_value(item.get("suggested_action")),
        ])

    filename = f"opportunities_{payload['compare_days']}d_{date.today().isoformat()}.csv"
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers=headers)


@app.get("/api/inventory-weekly/overview")
def api_inventory_weekly_overview(
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    keyword: str = Query(default=""),
) -> dict:
    filters = build_filters(start_date=start_date, end_date=end_date, site=site, store=store, keyword=keyword)
    return dashboard_service.get_inventory_weekly_overview(filters)


@app.get("/api/inventory-weekly/trends")
def api_inventory_weekly_trends(
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    keyword: str = Query(default=""),
) -> dict:
    filters = build_filters(start_date=start_date, end_date=end_date, site=site, store=store, keyword=keyword)
    return dashboard_service.get_inventory_weekly_trends(filters)


@app.get("/api/inventory-weekly/details")
def api_inventory_weekly_details(
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    keyword: str = Query(default=""),
    warning_status: str = Query(default="all"),
    metric: str = Query(default="all"),
    sort_field: str = Query(default=""),
    sort_dir: str = Query(default=""),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=10, le=100),
) -> dict:
    filters = build_filters(start_date=start_date, end_date=end_date, site=site, store=store, keyword=keyword)
    return dashboard_service.get_inventory_weekly_details(
        filters,
        warning_status=warning_status,
        metric=metric,
        sort_field=sort_field,
        sort_dir=sort_dir,
        page=page,
        page_size=page_size,
    )


@app.get("/api/inventory-weekly/export")
def api_inventory_weekly_export(
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    keyword: str = Query(default=""),
    warning_status: str = Query(default="all"),
    metric: str = Query(default="all"),
) -> StreamingResponse:
    filters = build_filters(start_date=start_date, end_date=end_date, site=site, store=store, keyword=keyword)
    payload = dashboard_service.get_inventory_weekly_export_payload(
        filters,
        warning_status=warning_status,
        metric=metric,
    )
    output = io.StringIO(newline="")
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow([
        "周开始", "周结束", "快照日期", "站点", "店铺", "MSKU", "可用数量", "可用成本",
        "在途数量", "在途成本", "在仓数量", "在仓成本", "采购数量", "采购成本", "是否预警", "预警指标",
    ])
    metric_labels = {"available": "可用", "transit": "在途", "warehouse": "在仓", "plan": "采购"}
    for item in payload["items"]:
        writer.writerow([
            csv_cell_value(item.get("week_start")),
            csv_cell_value(item.get("week_end")),
            csv_cell_value(item.get("snapshot_date")),
            csv_cell_value(item.get("site")),
            csv_cell_value(item.get("store")),
            csv_cell_value(item.get("msku")),
            csv_cell_value(item.get("available_quantity")),
            csv_cell_value(item.get("available_cost")),
            csv_cell_value(item.get("transit_quantity")),
            csv_cell_value(item.get("transit_cost")),
            csv_cell_value(item.get("warehouse_quantity")),
            csv_cell_value(item.get("warehouse_cost")),
            csv_cell_value(item.get("plan_quantity")),
            csv_cell_value(item.get("plan_cost")),
            "是" if item.get("warning") else "否",
            "、".join(metric_labels.get(key, key) for key in item.get("warning_metrics", [])),
        ])
    filename = f"inventory_weekly_{date.today().isoformat()}.csv"
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers=headers)


@app.get("/api/replenishment")
def api_replenishment(
    snapshot_date: str = Query(default=""),
    level: str = Query(default="all"),
    category: str = Query(default="all"),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    keyword: str = Query(default=""),
    moq_status: str = Query(default="all"),
    category_period_days: int = Query(default=30),
    sort_field: str = Query(default=""),
    sort_dir: str = Query(default=""),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=10, le=100),
) -> dict:
    return replenishment_service.get_payload(
        snapshot_date=snapshot_date,
        level=level,
        category=category,
        site=site,
        store=store,
        keyword=keyword,
        moq_status=moq_status,
        category_period_days=category_period_days,
        sort_field=sort_field,
        sort_dir=sort_dir,
        page=page,
        page_size=page_size,
    )


@app.get("/api/replenishment/export")
def api_replenishment_export(
    snapshot_date: str = Query(default=""),
    level: str = Query(default="all"),
    category: str = Query(default="all"),
    category_period_days: int = Query(default=30),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    keyword: str = Query(default=""),
    moq_status: str = Query(default="all"),
    sort_field: str = Query(default=""),
    sort_dir: str = Query(default=""),
) -> StreamingResponse:
    payload = replenishment_service.get_export_payload(
        snapshot_date=snapshot_date,
        level=level,
        category=category,
        category_period_days=category_period_days,
        site=site,
        store=store,
        keyword=keyword,
        moq_status=moq_status,
        sort_field=sort_field,
        sort_dir=sort_dir,
    )
    disabled_stores = replenishment_service.get_non_replenishment_stores()
    content = build_replenishment_xlsx(payload, disabled_stores)
    filename = f"replenishment_{payload.get('snapshot_date') or snapshot_date or date.today().isoformat()}.xlsx"
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@app.get("/api/replenishment/country-metrics")
def api_replenishment_country_metrics(
    snapshot_date: str = Query(default=""),
    site: str = Query(default=""),
    store: str = Query(default=""),
    msku: str = Query(default=""),
    period_days: int = Query(default=30),
    sort_field: str = Query(default=""),
    sort_dir: str = Query(default=""),
) -> dict:
    return replenishment_service.get_country_metrics(
        snapshot_date=snapshot_date,
        site=site,
        store=store,
        msku=msku,
        period_days=period_days,
        sort_field=sort_field,
        sort_dir=sort_dir,
    )


@app.get("/api/replenishment/level-flow")
def api_replenishment_level_flow(
    snapshot_date: str = Query(default=""),
    level: str = Query(default="all"),
    flow_type: str = Query(default="all"),
    category: str = Query(default="all"),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    keyword: str = Query(default=""),
    category_period_days: int = Query(default=30),
) -> dict:
    return replenishment_service.get_level_flow(
        snapshot_date=snapshot_date,
        level=level,
        flow_type=flow_type,
        category=category,
        site=site,
        store=store,
        keyword=keyword,
        category_period_days=category_period_days,
    )


@app.get("/api/replenishment-tracking")
def api_replenishment_tracking(
    snapshot_date: str = Query(default=""),
    tracking_window_days: int = Query(default=7),
    level: str = Query(default="all"),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    keyword: str = Query(default=""),
    status: str = Query(default="all"),
    sort_field: str = Query(default=""),
    sort_dir: str = Query(default=""),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=10, le=100),
    category_period_days: int = Query(default=30),
) -> dict:
    return replenishment_tracking_service.get_payload(
        snapshot_date=snapshot_date,
        tracking_window_days=tracking_window_days,
        level=level,
        site=site,
        store=store,
        keyword=keyword,
        status=status,
        sort_field=sort_field,
        sort_dir=sort_dir,
        page=page,
        page_size=page_size,
        category_period_days=category_period_days,
    )


@app.get("/api/replenishment-tracking/detail")
def api_replenishment_tracking_detail(
    snapshot_date: str = Query(default=""),
    tracking_window_days: int = Query(default=30),
    site: str = Query(default=""),
    store: str = Query(default=""),
    msku: str = Query(default=""),
) -> dict:
    return replenishment_tracking_service.get_detail(
        snapshot_date=snapshot_date,
        tracking_window_days=tracking_window_days,
        site=site,
        store=store,
        msku=msku,
    )


@app.get("/api/replenishment-tracking-summary")
def api_replenishment_tracking_summary(
    cutoff_date: str = Query(default=""),
    entry_batch_days: int = Query(default=90),
    level: str = Query(default="all"),
    purchase_status: str = Query(default="all"),
    fba_status: str = Query(default="all"),
    summary_stage: str = Query(default="all"),
    level_flow_stage: str = Query(default="all"),
    detail_stage: str = Query(default=""),
    category_period_days: int = Query(default=30),
    history_level: str = Query(default="all"),
    product_category: str = Query(default="all"),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    keyword: str = Query(default=""),
    order_keyword: str = Query(default=""),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=10, le=100),
) -> dict:
    return replenishment_tracking_summary_service.get_payload(
        cutoff_date=cutoff_date,
        entry_batch_days=entry_batch_days,
        level=level,
        purchase_status=purchase_status,
        fba_status=fba_status,
        summary_stage=summary_stage,
        level_flow_stage=level_flow_stage,
        detail_stage=detail_stage,
        category_period_days=category_period_days,
        history_level=history_level,
        product_category=product_category,
        site=site,
        store=store,
        keyword=keyword,
        order_keyword=order_keyword,
        page=page,
        page_size=page_size,
    )


@app.get("/api/replenishment-tracking-summary/detail")
def api_replenishment_tracking_summary_detail(
    cutoff_date: str = Query(default=""),
    site: str = Query(default=""),
    store: str = Query(default=""),
    msku: str = Query(default=""),
) -> dict:
    return replenishment_tracking_summary_service.get_detail(
        cutoff_date=cutoff_date,
        site=site,
        store=store,
        msku=msku,
    )


@app.get("/api/return-goods")
def api_return_goods(
    snapshot_date: str = Query(default=""),
    period_days: int = Query(default=1, ge=1, le=30),
    country_category: str = Query(default="all"),
    seller_name_new: str = Query(default="all"),
    keyword: str = Query(default=""),
    stage: str = Query(default="all"),
    warning_type: str = Query(default="all"),
    quick_filter: str = Query(default="all"),
    inventory_status_filter: str = Query(default="all"),
    return_day: int = Query(default=0, ge=0, le=21),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=0, le=500),
) -> dict:
    safe_return_day = return_day if isinstance(return_day, int) else 0
    params = {
        "snapshot_date": snapshot_date,
        "period_days": period_days,
        "country_category": country_category,
        "seller_name_new": seller_name_new,
        "keyword": keyword,
        "stage": stage,
        "warning_type": warning_type,
        "quick_filter": quick_filter,
        "inventory_status_filter": inventory_status_filter,
        "page": page,
        "page_size": page_size,
    }
    if safe_return_day:
        params["return_day"] = safe_return_day
    return return_goods_service.get_payload(**params)


@app.get("/api/return-goods/detail")
def api_return_goods_detail(
    snapshot_date: str = Query(default=""),
    return_event_id: str = Query(default=""),
) -> dict:
    return return_goods_service.get_detail(
        snapshot_date=snapshot_date,
        return_event_id=return_event_id,
    )


@app.get("/api/return-goods/stage-detail")
def api_return_goods_stage_detail(
    snapshot_date: str = Query(default=""),
    period_days: int = Query(default=1, ge=1, le=30),
    country_category: str = Query(default="all"),
    seller_name_new: str = Query(default="all"),
    keyword: str = Query(default=""),
    stage_key: str = Query(default="observe"),
) -> dict:
    return return_goods_service.get_stage_detail(
        snapshot_date=snapshot_date,
        period_days=period_days,
        country_category=country_category,
        seller_name_new=seller_name_new,
        keyword=keyword,
        stage_key=stage_key,
    )


@app.get("/api/detail")
def api_detail(
    request: Request,
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    over_limit: str = Query(default="all"),
    daily_sales_band: str = Query(default="all"),
    margin_band: str = Query(default="all"),
    keyword: str = Query(default=""),
    sort_field: str = Query(default=""),
    sort_dir: str = Query(default=""),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=15, ge=1, le=100),
) -> dict:
    filters = build_filters(
        start_date=start_date,
        end_date=end_date,
        site=site,
        store=store,
        over_limit=over_limit,
        daily_sales_band=daily_sales_band,
        margin_band=margin_band,
        keyword=keyword,
    )
    filters["column_filters"] = price_review_column_filters(request)
    return dashboard_service.get_detail_payload(
        filters,
        page=page,
        page_size=page_size,
        sort_field=sort_field,
        sort_dir=sort_dir,
    )


@app.get("/api/detail/export")
def api_detail_export(
    request: Request,
    start_date: Optional[str] = Query(default=None),
    end_date: Optional[str] = Query(default=None),
    site: str = Query(default="all"),
    store: str = Query(default="all"),
    over_limit: str = Query(default="all"),
    daily_sales_band: str = Query(default="all"),
    margin_band: str = Query(default="all"),
    keyword: str = Query(default=""),
) -> StreamingResponse:
    filters = build_filters(
        start_date=start_date,
        end_date=end_date,
        site=site,
        store=store,
        over_limit=over_limit,
        daily_sales_band=daily_sales_band,
        margin_band=margin_band,
        keyword=keyword,
    )
    filters["column_filters"] = price_review_column_filters(request)
    payload = dashboard_service.get_detail_export_payload(filters)
    columns = payload["columns"]

    output = io.StringIO(newline="")
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow([csv_header_value(column) for column in columns])
    for row in payload["rows"]:
        writer.writerow([csv_cell_value(row.get(column), column) for column in columns])

    filename = f"product_detail_raw_{payload['start_date']}_{payload['end_date']}.csv"
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
    }
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers=headers)


@app.get("/api/detail/{item_id}")
def api_detail_record(
    item_id: str,
    trend_days: int = Query(default=30, ge=7, le=30),
) -> dict:
    return dashboard_service.get_detail_record(item_id, trend_days)


@app.get("/price-review", response_class=HTMLResponse)
def price_review_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "price_review.html",
        {"page": "price_review", "title": "调价效果复盘"},
    )


@app.get("/price-adjustments", response_class=HTMLResponse)
def price_adjustments_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request,
        "price_adjustments.html",
        {"page": "price_adjustments", "title": "调价日历"},
    )


@app.get("/api/price-adjustments/daily-counts")
def api_price_adjustments_daily_counts(
    days: int = Query(default=30, ge=7, le=90),
    include_all: bool = Query(default=False),
) -> dict:
    return {
        "items": price_review_service.get_daily_adjustment_counts(
            days=days,
            include_all=include_all,
        )
    }


@app.put("/api/price-adjustments/daily-notes/{adjust_date}")
def api_price_adjustments_daily_note(
    adjust_date: str,
    payload: AdjustmentDayNotePayload,
) -> dict:
    parsed_date = _parse_date(adjust_date)
    if parsed_date is None:
        raise HTTPException(status_code=400, detail="Invalid adjust_date")
    try:
        note = price_review_service.save_adjustment_day_note(parsed_date, payload.note)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"date": parsed_date.isoformat(), "note": note}


def price_review_filters(
    adjust_date: Optional[str] = Query(default=None),
    compare_days: int = Query(default=3, ge=3, le=28),
    country: str = Query(default=""),
    drop_range: str = Query(default=""),
    risk_level: str = Query(default=""),
    store: str = Query(default=""),
    price_band: str = Query(default=""),
    adjustment_type: str = Query(default=""),
    keyword: str = Query(default=""),
) -> dict:
    return {
        "adjust_date": _parse_date(adjust_date),
        "compare_days": compare_days,
        "country": country,
        "drop_range": drop_range,
        "risk_level": risk_level,
        "store": store,
        "price_band": price_band,
        "adjustment_type": adjustment_type,
        "keyword": keyword,
    }


def price_review_column_filters(request: Request) -> dict[str, str]:
    return {
        key[3:]: value
        for key, value in request.query_params.items()
        if key.startswith("cf_") and value not in {"", "all"}
    }


@app.get("/api/price-review/overview")
def api_price_review_overview(filters: dict = Depends(price_review_filters)) -> dict:
    return price_review_service.get_overview_payload(**filters)


@app.get("/api/price-review/drop-range")
def api_price_review_drop_range(filters: dict = Depends(price_review_filters)) -> dict:
    return price_review_service.get_drop_range_payload(**filters)


@app.get("/api/price-review/matrices")
def api_price_review_matrices(filters: dict = Depends(price_review_filters)) -> dict:
    return price_review_service.get_matrices_payload(**filters)


@app.get("/api/price-review/countries")
def api_price_review_countries(filters: dict = Depends(price_review_filters)) -> dict:
    return price_review_service.get_country_payload(**filters)


@app.get("/api/price-review/second-adjustments")
def api_price_review_second_adjustments(filters: dict = Depends(price_review_filters)) -> dict:
    return price_review_service.get_second_adjustments_payload(**filters)


@app.get("/api/price-review/top-lists")
def api_price_review_top_lists(filters: dict = Depends(price_review_filters)) -> dict:
    return price_review_service.get_top_lists_payload(**filters)


@app.get("/api/price-review/skus")
def api_price_review_skus(
    request: Request,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    sort_field: str = Query(default=""),
    sort_dir: str = Query(default=""),
    filters: dict = Depends(price_review_filters),
) -> dict:
    filters = dict(filters)
    filters["column_filters"] = price_review_column_filters(request)
    return price_review_service.get_sku_list_payload(
        page=page,
        page_size=page_size,
        sort_field=sort_field,
        sort_dir=sort_dir,
        **filters,
    )


@app.get("/api/price-review/role-migration")
def api_price_review_role_migration(
    adjust_date: Optional[str] = Query(default=None),
    pre_days: Optional[int] = Query(default=None),
    post_days: int = Query(default=3),
    comparison_mode: Optional[str] = Query(default=None),
    country: str = Query(default=""),
    store: str = Query(default=""),
    role_before: str = Query(default=""),
    role_after: str = Query(default=""),
    role_change: str = Query(default=""),
    finance_change: str = Query(default=""),
    finance_before: str = Query(default=""),
    finance_after: str = Query(default=""),
    data_status: str = Query(default=""),
    keyword: str = Query(default=""),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    sort_field: str = Query(default=""),
    sort_dir: str = Query(default=""),
    column_filters: str = Query(default=""),
) -> dict:
    # Compatibility for previously shared links. New clients always send
    # pre_days explicitly; old equal-window links map N to both sides.
    resolved_pre_days = pre_days
    if resolved_pre_days is None:
        resolved_pre_days = post_days if comparison_mode == "equal_window" else 30
    return price_review_service.get_role_migration_payload(
        adjust_date=_parse_date(adjust_date),
        pre_days=resolved_pre_days,
        post_days=post_days,
        country=country,
        store=store,
        role_before=role_before,
        role_after=role_after,
        role_change_filter=role_change,
        finance_change_filter=finance_change,
        finance_before=finance_before,
        finance_after=finance_after,
        data_status=data_status,
        keyword=keyword,
        page=page,
        page_size=page_size,
        sort_field=sort_field,
        sort_dir=sort_dir,
        column_filters=column_filters,
    )


@app.get("/api/price-review/role-migration/detail")
def api_price_review_role_migration_detail(
    adjust_date: str = Query(...),
    pre_days: int = Query(default=30),
    post_days: int = Query(default=3),
    country: str = Query(...),
    store: str = Query(...),
    msku: str = Query(...),
) -> dict:
    parsed_date = _parse_date(adjust_date)
    if parsed_date is None:
        raise HTTPException(status_code=400, detail="调价日期格式无效")
    payload = price_review_service.get_role_migration_detail(
        adjust_date=parsed_date,
        pre_days=pre_days,
        post_days=post_days,
        country=country,
        store=store,
        msku=msku,
    )
    if payload is None:
        raise HTTPException(status_code=404, detail="未找到该站点 SKU 的角色迁移记录")
    return payload


_TOP_LIST_COLUMNS = [
    ("类型", "类型"),
    ("国家", "country"),
    ("店铺", "store"),
    ("MSKU", "msku"),
    ("品名", "product_name"),
    ("调价类型", "adjustment_type"),
    ("上次调价日期", "previous_adjust_date"),
    ("前销量", "sales_before"),
    ("后销量", "sales_after"),
    ("销量变化", "sales_change"),
    ("前日销", "daily_sales_before"),
    ("后日销", "daily_sales_after"),
    ("日销变化", "daily_sales_change"),
    ("毛利润变化", "profit_change"),
    ("后毛利率", "margin_after"),
    ("排名变化", "rank_change"),
    ("风险等级", "risk_level"),
    ("问题标签", "issue_tags"),
]

_SKU_EXPORT_COLUMNS = [
    ("国家", "country"),
    ("店铺", "store"),
    ("MSKU", "msku"),
    ("品名", "product_name"),
    ("调价类型", "adjustment_type"),
    ("上次调价日期", "previous_adjust_date"),
    ("调价前价格", "price_before"),
    ("调价后价格", "price_after"),
    ("调价幅度", "drop_ratio"),
    ("调价幅度区间", "drop_range"),
    ("价格带", "price_band"),
    ("前销量", "sales_before"),
    ("后销量", "sales_after"),
    ("销量变化", "sales_change"),
    ("前日销", "daily_sales_before"),
    ("后日销", "daily_sales_after"),
    ("前销售额", "revenue_before"),
    ("后销售额", "revenue_after"),
    ("前毛利润", "profit_before"),
    ("后毛利润", "profit_after"),
    ("前毛利率", "margin_before"),
    ("后毛利率", "margin_after"),
    ("毛利率变化", "margin_change"),
    ("前Sessions", "sessions_before"),
    ("后Sessions", "sessions_after"),
    ("前ACOS", "acos_before"),
    ("后ACOS", "acos_after"),
    ("前TACOS", "tacos_before"),
    ("后TACOS", "tacos_after"),
    ("前CPC", "cpc_before"),
    ("后CPC", "cpc_after"),
    ("调前排名", "rank_before"),
    ("调后排名", "rank_after"),
    ("排名变化", "rank_change"),
    ("调前销量分层", "sales_band_before"),
    ("调后销量分层", "sales_band_after"),
    ("调前日销分层", "daily_sales_band_before"),
    ("调后日销分层", "daily_sales_band_after"),
    ("调前毛利率分层", "margin_band_before"),
    ("调后毛利率分层", "margin_band_after"),
    ("调前排名分层", "rank_band_before"),
    ("调后排名分层", "rank_band_after"),
    ("风险等级", "risk_level"),
    ("问题标签", "issue_tags"),
    ("建议动作", "suggested_action"),
]


def _write_price_review_csv(rows: list[dict], columns: list[tuple[str, str]], filename: str) -> StreamingResponse:
    output = io.StringIO(newline="")
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow([header for header, _ in columns])
    for row in rows:
        writer.writerow([csv_cell_value(row.get(key)) for _, key in columns])
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
    }
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers=headers)


@app.get("/api/price-review/top-lists/export")
def api_price_review_top_lists_export(
    filters: dict = Depends(price_review_filters),
) -> StreamingResponse:
    rows = price_review_service.get_top_lists_export_payload(**filters)
    date_label = filters["adjust_date"].isoformat() if filters.get("adjust_date") else "latest"
    filename = f"price_review_top_lists_{date_label}_{filters['compare_days']}d.csv"
    return _write_price_review_csv(rows, _TOP_LIST_COLUMNS, filename)


@app.get("/api/price-review/skus/export")
def api_price_review_skus_export(
    request: Request,
    filters: dict = Depends(price_review_filters),
) -> StreamingResponse:
    filters = dict(filters)
    filters["column_filters"] = price_review_column_filters(request)
    rows = price_review_service.get_sku_list_export_payload(**filters)
    date_label = filters["adjust_date"].isoformat() if filters.get("adjust_date") else "latest"
    filename = f"price_review_skus_{date_label}_{filters['compare_days']}d.csv"
    return _write_price_review_csv(rows, _SKU_EXPORT_COLUMNS, filename)
